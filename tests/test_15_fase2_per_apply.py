from __future__ import annotations

from io import BytesIO
from pathlib import Path
import sys
import tempfile
from unittest.mock import patch
from zipfile import ZipFile

import base64
import pandas as pd
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]

if str(ROOT) not in sys.path:
    sys.path.insert(
        0,
        str(ROOT),
    )


from app.services.provider_excel_remote_status_service import (
    ProviderExcelRemoteStatusService,
)
from app.services.provider_excel_remote_apply_service import (
    ProviderExcelRemoteApplyService,
    ProviderExcelRemoteLockedError,
)


errors = []


def ok(message):
    print(
        f"[OK] {message}"
    )


def fail(message):
    print(
        f"[ERROR] {message}"
    )
    errors.append(
        message
    )


def build_cached_formula_xlsx() -> bytes:
    encoded = (
        "UEsDBBQAAAAIAAAAIQBhXUk6TwEAAI8EAAATAAAAW0NvbnRlbnRfVHlwZXNdLnhtbK2Uy27CMBBF9/2KyNsqMXRRVRWBRR/LFqn0A1x7Qiwc2/IMFP6+k/BQW1Gggk2sZO7cc8eOPBgtG5ctIKENvhT9oicy8DoY66eleJ8853ciQ1LeKBc8lGIFKEbDq8FkFQEzbvZYipoo3kuJuoZGYREieK5UITWK+DVNZVR6pqYgb3q9W6mDJ/CUU+shhoNHqNTcUfa05M/rIAkciuxhLWxZpVAxOqsVcV0uvPlFyTeEgjs7DdY24jULhNxLaCt/AzZ9r7wzyRrIxirRi2pYJU3Q4xQiStYXh132xAxVZTWwx7zhlgLaQAZMHtkSElnYZT7I1iHB/+HbPWq7TyQunURaOcCzR8WYQBmsAahxxdr0CJn4f4L1s382v7M5AvwMafYRwuzSw7Zr0SjrT+B3YpTdcv7UP4Ps/I8dea0SmDdKfA1c/OS/e29zyO4+GX4BUEsDBBQAAAAIAAAAIQDyn0na6QAAAEsCAAALAAAAX3JlbHMvLnJlbHOtksFOwzAMQO98ReT7mm5ICKGluyCk3SY0PsAkbhu1jaPEg+7viZBADI1pB45x7Odny+vNPI3qjVL2HAwsqxoUBcvOh87Ay/5pcQ8qCwaHIwcycKQMm+Zm/UwjSqnJvY9ZFUjIBnqR+KB1tj1NmCuOFMpPy2lCKc/U6Yh2wI70qq7vdPrJgOaEqbbOQNq6Jaj9MdI1bG5bb+mR7WGiIGda/MooZEwdiYF51O+chlfmoSpQ0OddVte7/D2nnkjQoaC2nGgRU6lO4stav3Uc210J58+MS0K3/7kcmoWCI3dZCWP8MtInN9B8AFBLAwQUAAAACAAAACEARHVb8OgAAAC5AgAAGgAAAHhsL19yZWxzL3dvcmtib29rLnhtbC5yZWxzrZLBasMwEETv/Qqx91p2EkopkXMphVzb9AOEtLZMbElot2n99xEJTR0IoQefxIzYmQe7683P0IsDJuqCV1AVJQj0JtjOtwo+d2+PzyCItbe6Dx4VjEiwqR/W79hrzjPkukgih3hS4Jjji5RkHA6aihDR558mpEFzlqmVUZu9blEuyvJJpmkG1FeZYmsVpK2tQOzGiP/JDk3TGXwN5mtAzzcq5HdIe3KInEN1apEVXCySp6cqcirI2zCLOWE4z+IfyEmezbsMyzkZiMc+L/QCcdb36lez1jud0H5wytc2pZjavzDy6uLqI1BLAwQUAAAACAAAACEAXETGxuoBAABTBQAAGAAAAHhsL3dvcmtzaGVldHMvc2hlZXQxLnhtbI2UW2/bIBSA3/crEK/Tgu0svUS2q7TOreqUKusq7ZHYxxfVBgto0v37gdNEwPKwN+DjHA7fEcR3H12L9iBkw1mCw1GAEbCcFw2rEvzrZfHtBiOpKCtoyxkk+A9IfJd+iQ9cvMkaQCGdgMkE10r1U0JkXkNH5Yj3wDQpueio0lNREdkLoMUQ1LUkCoIr0tGG4WOGqfifHLwsmxwynr93wNQxiYCWKl2+rJte4jQuGs3MfZCAMsGzcPo4xiSNh5NfGzhIa4wU3f2EFnIFhb4/RuZiO87fDFzrpcCEkn9iF0NRzwIVUNL3Vm35YQVNVSudZHI+LaOKprHgBySG5LKnxlU4DQNdaG5WZ3pZB5nC92kQk70+Lf9k9zYLXfZgs8hlmc3GLpvb7LvLFjabuGxpsyuXrWx27bK1zW5c9miz2zMjWtjZWnTZWmSb8bU50POWfUIlNC7Tr5nmpdmYbZ6eZtuH9ctvT5iTzTO9cKCneulAz/XKq2N1quN5u3mdz7PN1vPoJJtcljW+LGtsh3q9u3eg17zMgV735g689cTYMPIatHSg16CVAz3fawd6vvUzdx/R0Q2xnmJPK/hBRdUwiVoodapgdI2ROL7cYax4P4wmGO24Urw7zWr9e4EwM11Eybk6TcwHcf4P079QSwMEFAAAAAgAAAAhADBPQT5MAQAAKwIAAA8AAAB4bC93b3JrYm9vay54bWyNUctOwzAQvPMV1t5pHkoiWjWpxEtUQtADtGcTbxqrjh3ZTtP+PZtUKXDj5J3ZndHOerk6NYod0TppdA7RLASGujRC6n0Onx/Pt3fAnOdacGU05nBGB6viZtkbe/gy5sBIr10OtfftIghcWWPD3cy0qKlTGdtwT9DuA9da5MLViL5RQRyGWdBwqeHisLD/8TBVJUt8NGXXoPYXE4uKe9re1bJ1UCwrqXB7CcR4277xhtY+KWCKO/8kpEeRQ0LQ9PiHsF1730k1gDRMISiuITeWCax4p/wHrTa507niJI6zYXKY2krs3Y9ogOy0k1qYPoc4ocueJxSlwPqx3knhayKycH7lXlDua5/DPMvCwTz45T7eb3qZHsNtrDkiCmPpb9jYWFOIiBItJBV2LaLRZtKWXJUUaXjGwThJozmwqlPqgbh3/Wr4aDCIpjjFN1BLAwQUAAAACAAAACEAyGtidTgBAAD9AgAAFAAAAHhsL3NoYXJlZFN0cmluZ3MueG1sdZLda8IwFMXf91eEvM/U+slo6/zoQJBVig62t5jcaaDN7ZJUtv9+FVFGs+Yt55dz7uGSaPZdFuQMxirUMe33AkpAC5RKH2O63708TimxjmvJC9QQ0x+wdJY8RNY60li1jenJueqJMStOUHLbwwp0Qz7RlNw1V3NktjLApT0BuLJgYRCMWcmVpkRgrV1MwyEltVZfNSzvQhJZlUQu0VgeDETMJRG7KFeVV1AUSmKX3owGoz18DVsWCrTzMk0ttgbPABLN/8ZOLNAY8KY5VeHe1twoD0kUa9mUUJJLL+yyLM8xf523pW2apx9tMQz6t+O9z7O3NF1lOfGi+LMD63oCyzbph4PhaDyZtvVFuvNSNtnWL7TKNpt5vlzv3v2q4e10V1200aGz6t3UBtPJeDQchH8Wwpr/m/wCUEsDBBQAAAAIAAAAIQAavBrJjwEAAGMDAAANAAAAeGwvc3R5bGVzLnhtbKWTUWvcMAzH3/cpjN/b3B2srCNJHwoHhW0MeoW9+mIlZ7DlYCvHZZ9+cpwmdzDoQ58s/yX9JFt2+XRxVpwhROOxktv7jRSAjdcGu0q+HfZ336SIpFAr6xEqOUKUT/WXMtJo4fUEQIIJGCt5Iuq/F0VsTuBUvPc9IHtaH5wi3oauiH0ApWNKcrbYbTYPhVMGZV22HimKxg9I3MQs1GX8K87KsrKVRV023vogiPGQglhB5SBHPCtrjsEksVXO2DHLuyRMHc1xzqAPSSxyhWmJnGSsXRrYySzUZa+IIOCeN2K2D2PP5ZEvI2OmuA+iu6DG7e7rVcK0cN2jD5ov//roWapLCy1xQjDdKa3k+yI5ibxjQxvVeVQ2Id8zZoOxDVj7mib0p71hX1qBg9s7etGV5FGn07+b3NBsZkzeJP41LbM/jRWX9pa/oKdCN/RFFWnelfyVnpRdEeI4GEsG/9MwM/Vl7XXykjryy72pwgwNrRosHRZnJVf7J2gzuMcl6rc5e5qjVvtHmtT2Yepg/R71P1BLAwQUAAAACAAAACEAGPpGVLAFAABSGwAAEwAAAHhsL3RoZW1lL3RoZW1lMS54bWztWU2P20QYvvMrRr63jhM7za6arTbZpIXttqvdtKjHiT2xpxl7rJnJbnND7REJCVEQFyRuHBBQqZW4lF+zUARF6l/g9UeS8WayzbaLALU5JJ7x835/+B3n6rUHMUNHREjKk7blXK5ZiCQ+D2gStq07g/6lloWkwkmAGU9I25oSaV3b+uAq3lQRiQkC8kRu4rYVKZVu2rb0YRvLyzwlCdwbcRFjBUsR2oHAx8A2Zna9VmvaMaaJhRIcA9fboxH1CRpkLK2tGfMeg69EyWzDZ+LQzyXqFDk2GDvZj5zKLhPoCLO2BXICfjwgD5SFGJYKbrStWv6x7K2r9pyIqRW0Gl0//5R0JUEwrud0IhzOCZ2+u3FlZ86/XvBfxvV6vW7PmfPLAdj3wVJnCev2W05nxlMDFZfLvLs1r+ZW8Rr/xhJ+o9PpeBsVfGOBd5fwrVrT3a5X8O4C7y3r39nudpsVvLfAN5fw/SsbTbeKz0ERo8l4CZ3Fcx6ZOWTE2Q0jvAXw1iwBFihby66CPlGrci3G97noAyAPLlY0QWqakhH2AdfF8VBQnAnAmwRrd4otXy5tZbKQ9AVNVdv6KMVQEQvIq+c/vHr+FL16/uTk4bOThz+fPHp08vAnA+ENnIQ64cvvPv/rm0/Qn0+/ffn4SzNe6vjffvz011++MAOVDnzx1ZPfnz158fVnf3z/2ADfFniowwc0JhLdIsfogMdgm0EAGYrzUQwiTCsUOAKkAdhTUQV4a4qZCdchVefdFdAATMDrk/sVXQ8jMVHUANyN4gpwj3PW4cJozm4mSzdnkoRm4WKi4w4wPjLJ7p4KbW+SQiZTE8tuRCpq7jOINg5JQhTK7vExIQaye5RW/LpHfcElHyl0j6IOpkaXDOhQmYlu0BjiMjUpCKGu+GbvLupwZmK/Q46qSCgIzEwsCau48TqeKBwbNcYx05E3sYpMSh5OhV9xuFQQ6ZAwjnoBkdJEc1tMK+ruYuhExrDvsWlcRQpFxybkTcy5jtzh426E49SoM00iHfuhHEOKYrTPlVEJXq2QbA1xwMnKcN+lRJ2vrO/QMDInSHZnIsquXem/MU3OasaMQjd+34xn8G14NJlK4nQLXoX7HzbeHTxJ9gnk+vu++77vvot9d1Utr9ttFw3W1ufinF+8ckgeUcYO1ZSRmzJvzRKUDvqwmS9yovlMnkZwWYqr4EKB82skuPqYqugwwimIcXIJoSxZhxKlXMJJwFrJOz9OUjA+3/NmZ0BAY7XHg2K7oZ8N52zyVSh1QY2MwbrCGlfeTphTANeU5nhmad6Z0mzNm1ANCGcHf6dZL0RDxmBGgszvBYNZWC48RDLCASlj5BgNcRpruq31eq9p0jYabydtnSDp4twV4rwLiFJtKUr2cjmypLpCx6CVV/cs5OO0bY1gkoLLOAV+MmtAmIVJ2/JVacpri/m0wea0dGorDa6ISIVUO1hGBVV+a/bqJFnoX/fczA8XY4ChG62nRaPl/Ita2KdDS0Yj4qsVO4tleY9PFBGHUXCMhmwiDjDo7RbZFVAJz4z6bCGgQt0y8aqVX1bB6Vc0ZXVglka47EktLfYFPL+e65CvNPXsFbq/oSmNCzTFe3dNyTIXxtZGkB+oYAwQGGU52ra4UBGHLpRG1O8LGBxyWaAXgrLIVEIse9+c6UqOFn2r4FE0uTBSBzREgkKnU5EgZF+Vdr6GmVPXn68zRmWfmasr0+J3SI4IG2TV28zst1A06yalI3Lc6aDZpuoahv3/8OTjrph8zh4PFoLc88wirtb0tUfBxtupcM5Hbd1scd1b+1GbwuEDZV/QuKnw2WK+HfADiD6aT5QIEvFSqyy/+eYQdG5pxmWs/tkxahGC1op4X+TwqTm7scLZZ4t7c2d7Bl97Z7vaXi5RWzvI5KulP5748D7I3oGD0oQpWbxNegBHze7sLwPgYy9It/4GUEsDBBQAAAAIAAAAIQDCANFYJgEAAFACAAARAAAAZG9jUHJvcHMvY29yZS54bWydks1OwzAQhO88ReR74iSlVbGSVALUE5WQKAJxs+xtahH/yDakeXvcpE1bKSeO65n9dnblYnWQTfQL1gmtSpQlKYpAMc2Fqkv0vl3HSxQ5TxWnjVZQog4cWlV3BTOEaQuvVhuwXoCLAkg5wkyJ9t4bgrFje5DUJcGhgrjTVlIfSltjQ9k3rQHnabrAEjzl1FN8BMZmJKITkrMRaX5s0wM4w9CABOUdzpIMX7werHSTDb1y5ZTCdwYmrWdxdB+cGI1t2ybtrLeG/Bn+3Ly89avGQh1PxQBVBWeEWaBe26rA10U4XEOd34QT7wTwxy7oE2+nRYY+4FEIQIa4Z+Vj9vS8XaMqT/NFnD7E6f02W5J8Rubzr+PIm/4LUJ6G/Jt4Bgy5bz9B9QdQSwMEFAAAAAgAAAAhAC5SM8F7AQAAFQMAABAAAABkb2NQcm9wcy9hcHAueG1snVLBbtswDL3vKwzdGznBMAyBrKJIO/SwYQGStmdOpmOhsiSIrJHs6yc7iOusO02nR/Lh6emJ6vbYuaLHRDb4SiwXpSjQm1Bbf6jE0/7bzVdREIOvwQWPlTghiVv9SW1TiJjYIhVZwVMlWua4lpJMix3QIo99njQhdcC5TAcZmsYavA/mrUPPclWWXyQeGX2N9U2cBMVZcd3z/4rWwQz+6Hl/illPq7sYnTXA+ZH6hzUpUGi4eDgadErOhyoL7dC8JcsnXSo5L9XOgMNNFtYNOEIl3xvqEWHIbAs2kVY9r3s0HFJB9ndObSWKX0A42KlED8mCZ3GmnYsRu0ic9EtIr9QiMik5NUc4586x/ayXIyGDa6KcjGR8bXFv2SH9bLaQ+B+Ol3PHowcx85jT7RHrkPCjyct1f12wCV0En1OUE/pu/Ss9xX24B8ZLptdNtWshYZ2/Ycp8aqjHbC65gb9pwR+wvnA+DoYNeD5vuV6uFmU+48dfekq+L7T+A1BLAQIUAxQAAAAIAAAAIQBhXUk6TwEAAI8EAAATAAAAAAAAAAAAAACAAQAAAABbQ29udGVudF9UeXBlc10ueG1sUEsBAhQDFAAAAAgAAAAhAPKfSdrpAAAASwIAAAsAAAAAAAAAAAAAAIABgAEAAF9yZWxzLy5yZWxzUEsBAhQDFAAAAAgAAAAhAER1W/DoAAAAuQIAABoAAAAAAAAAAAAAAIABkgIAAHhsL19yZWxzL3dvcmtib29rLnhtbC5yZWxzUEsBAhQDFAAAAAgAAAAhAFxExsbqAQAAUwUAABgAAAAAAAAAAAAAAIABsgMAAHhsL3dvcmtzaGVldHMvc2hlZXQxLnhtbFBLAQIUAxQAAAAIAAAAIQAwT0E+TAEAACsCAAAPAAAAAAAAAAAAAACAAdIFAAB4bC93b3JrYm9vay54bWxQSwECFAMUAAAACAAAACEAyGtidTgBAAD9AgAAFAAAAAAAAAAAAAAAgAFLBwAAeGwvc2hhcmVkU3RyaW5ncy54bWxQSwECFAMUAAAACAAAACEAGrwayY8BAABjAwAADQAAAAAAAAAAAAAAgAG1CAAAeGwvc3R5bGVzLnhtbFBLAQIUAxQAAAAIAAAAIQAY+kZUsAUAAFIbAAATAAAAAAAAAAAAAACAAW8KAAB4bC90aGVtZS90aGVtZTEueG1sUEsBAhQDFAAAAAgAAAAhAMIA0VgmAQAAUAIAABEAAAAAAAAAAAAAAIABUBAAAGRvY1Byb3BzL2NvcmUueG1sUEsBAhQDFAAAAAgAAAAhAC5SM8F7AQAAFQMAABAAAAAAAAAAAAAAAIABpREAAGRvY1Byb3BzL2FwcC54bWxQSwUGAAAAAAoACgCAAgAAThMAAAAA"
    )

    return base64.b64decode(
        encoded
    )


def relation_results() -> pd.DataFrame:
    return pd.DataFrame([
        {
            "_item_id":
                "EXCEL-000001",
            "email":
                "a@test.com",
            "nombre":
                "ANA",
            "apellido_pat":
                "PEREZ",
            "apellido_mat":
                "",
            "nombre_proveedor":
                "PROVEEDOR A",
            "nombre_cliente":
                "DOLLARCITY",
            "creado_objetivo":
                1,
            "_origen_nombre":
                "ANA",
            "_origen_apellido":
                "PEREZ",
            "_origen_apellidomaterno":
                "",
            "_origen_nombreCliente":
                "DOLLARCITY",
            "_origen_rucProveedor":
                "20111111111",
            "_origen_nombreProveedor":
                "PROVEEDOR A",
            "_origen_correo":
                "a@test.com",
            "_origen_tipoUsuario":
                "PROVEEDOR",
            "_origen_docIdentidad":
                "12345678",
        },
        {
            "_item_id":
                "EXCEL-000002",
            "email":
                "b@test.com",
            "nombre":
                "BETA",
            "apellido_pat":
                "LOPEZ",
            "apellido_mat":
                "",
            "nombre_proveedor":
                "PROVEEDOR B",
            "nombre_cliente":
                "DOLLARCITY",
            "creado_objetivo":
                2,
            "_origen_nombre":
                "BETA",
            "_origen_apellido":
                "LOPEZ",
            "_origen_apellidomaterno":
                "",
            "_origen_nombreCliente":
                "DOLLARCITY",
            "_origen_rucProveedor":
                "20222222222",
            "_origen_nombreProveedor":
                "PROVEEDOR B",
            "_origen_correo":
                "b@test.com",
            "_origen_tipoUsuario":
                "PROVEEDOR",
            "_origen_docIdentidad":
                "87654321",
        },
    ])


class FakeStatusService(
    ProviderExcelRemoteStatusService
):
    def __init__(
        self,
        original: bytes,
    ):
        self.original = original
        self.uploaded = None

    def download(
        self,
        shared_url: str,
    ) -> bytes:
        return (
            self.uploaded
            if self.uploaded is not None
            else self.original
        )


class FakeApplyService(
    ProviderExcelRemoteApplyService
):
    def __init__(
        self,
        status_service,
    ):
        super().__init__(
            status_service=
                status_service
        )
        self.resolve_calls = 0
        self.upload_calls = 0

    def resolve_drive_item(
        self,
        shared_url: str,
    ) -> dict:
        self.resolve_calls += 1

        return {
            "item_id": "ITEM-1",
            "drive_id": "DRIVE-1",
            "name": "usuarios.xlsx",
            "etag": "ETAG-1",
            "size": 100,
        }

    def upload_content(
        self,
        metadata: dict,
        content: bytes,
    ) -> dict:
        self.upload_calls += 1
        self.status_service.uploaded = content

        return {
            "id": "ITEM-1",
            "eTag": "ETAG-2",
        }


class FakePutResponse:
    def __init__(
        self,
        status_code: int,
        text: str = "",
        headers: dict | None = None,
        payload: dict | None = None,
    ):
        self.status_code = status_code
        self.text = text
        self.headers = (
            headers
            or {}
        )
        self.ok = (
            200
            <= status_code
            < 300
        )
        self._payload = (
            payload
            or {}
        )

    def json(self):
        return self._payload

    def __enter__(self):
        return self

    def __exit__(
        self,
        exc_type,
        exc_value,
        traceback,
    ):
        return False


print("=" * 100)
print(
    "TEST 15 - APLICACION FASE 2 PER"
)
print("=" * 100)

original = build_cached_formula_xlsx()
relations = relation_results()

status_service = (
    ProviderExcelRemoteStatusService()
)

preview = status_service.build_preview(
    content=original,
    relation_results=relations,
)

if (
    preview["ready"]
    and preview["summary"][
        "actualizaciones"
    ] == 2
):
    ok(
        "Preview previo queda listo para escritura"
    )
else:
    fail(
        "Preview previo no quedo listo"
    )

apply_service = (
    ProviderExcelRemoteApplyService(
        status_service=status_service,
    )
)

updated = apply_service.build_updated_content(
    content=original,
    plan=preview["plan"],
)

formula_book = load_workbook(
    BytesIO(updated),
    data_only=False,
)

formula_sheet = formula_book[
    "Proveedores"
]

formula_ok = (
    formula_sheet["D2"].value
    == "=+D1"
    and formula_sheet["H2"].value
    == "=+H1"
    and formula_sheet["J2"].value
    == 1
    and formula_sheet["J3"].value
    == 2
)

formula_book.close()

if formula_ok:
    ok(
        "Actualizacion XML conserva formulas y cambia creado"
    )
else:
    fail(
        "La actualizacion XML altero formulas o estados"
    )

value_book = load_workbook(
    BytesIO(updated),
    data_only=True,
)

value_sheet = value_book[
    "Proveedores"
]

cached_ok = (
    value_sheet["D2"].value
    == "DOLLARCITY"
    and value_sheet["H2"].value
    == "PROVEEDOR"
)

value_book.close()

if cached_ok:
    ok(
        "Valores calculados en cache se conservan"
    )
else:
    fail(
        "Se perdieron valores calculados de formulas"
    )

with ZipFile(
    BytesIO(original),
    "r",
) as before_zip:
    with ZipFile(
        BytesIO(updated),
        "r",
    ) as after_zip:
        changed = [
            name
            for name in before_zip.namelist()
            if before_zip.read(name)
            != after_zip.read(name)
        ]

if (
    len(changed) == 1
    and changed[0].startswith(
        "xl/worksheets/"
    )
):
    ok(
        "Solo cambia el XML de la hoja Proveedores"
    )
else:
    fail(
        "Se modificaron componentes XLSX adicionales: "
        + str(changed)
    )

post_verify = status_service.build_preview(
    content=updated,
    relation_results=relations,
)

if (
    post_verify["ready"]
    and post_verify["summary"][
        "actualizaciones"
    ] == 0
    and post_verify["summary"][
        "sin_cambio"
    ] == 2
):
    ok(
        "Contenido actualizado pasa verificacion posterior"
    )
else:
    fail(
        "Contenido actualizado no pasa verificacion posterior"
    )

with tempfile.TemporaryDirectory() as temp_dir:
    fake_status = FakeStatusService(
        original
    )

    fake_apply = FakeApplyService(
        fake_status
    )

    result = fake_apply.apply(
        shared_url="https://example.invalid/share",
        relation_results=relations,
        output_dir=temp_dir,
    )

    if (
        result["ok"]
        and result["created_1"] == 1
        and result["created_2"] == 1
        and result["verified_rows"] == 2
        and fake_apply.upload_calls == 1
    ):
        ok(
            "Flujo apply sube una vez y verifica estados"
        )
    else:
        fail(
            "Flujo apply no produjo el resultado esperado"
        )

    paths_ok = all(
        Path(result[key]).is_file()
        for key in [
            "backup_path",
            "candidate_path",
            "verified_path",
            "report_path",
        ]
    )

    if paths_ok:
        ok(
            "Apply genera backup, objetivo, verificado y reporte"
        )
    else:
        fail(
            "Faltan archivos de auditoria del apply"
        )

# ------------------------------------------------------------------
# HTTP 423 - lock transitorio
# ------------------------------------------------------------------

retry_service = (
    ProviderExcelRemoteApplyService()
)

locked_body = (
    '{"error":{"code":"notAllowed",'
    '"message":"locked",'
    '"innerError":{"code":"resourceLocked"}}}'
)

responses = [
    FakePutResponse(
        423,
        text=locked_body,
    ),
    FakePutResponse(
        423,
        text=locked_body,
    ),
    FakePutResponse(
        200,
        payload={
            "id": "ITEM-1",
            "eTag": "ETAG-2",
        },
    ),
]

with patch("app.services.provider_excel_remote_apply_service.MsalAuthService") as auth_mock, patch(
    "app.services."
    "provider_excel_remote_apply_service."
    "requests.put",
    side_effect=responses,
) as put_mock:
    auth_mock.return_value.get_access_token.return_value = "test-token"
    with patch(
        "app.services."
        "provider_excel_remote_apply_service."
        "time.sleep"
    ):
        retry_result = (
            retry_service.upload_content(
                metadata={
                    "drive_id": "DRIVE-1",
                    "item_id": "ITEM-1",
                },
                content=b"xlsx-content",
            )
        )

if (
    retry_result.get("id")
    == "ITEM-1"
    and put_mock.call_count == 3
):
    ok(
        "HTTP 423 transitorio reintenta "
        "y luego completa upload"
    )
else:
    fail(
        "HTTP 423 transitorio no reintento "
        "correctamente"
    )


# ------------------------------------------------------------------
# HTTP 423 - lock persistente
# ------------------------------------------------------------------

persistent_responses = [
    FakePutResponse(
        423,
        text=locked_body,
    )
    for _ in range(
        retry_service.UPLOAD_MAX_RETRIES
    )
]

persistent_lock_ok = False

with patch("app.services.provider_excel_remote_apply_service.MsalAuthService") as auth_mock, patch(
    "app.services."
    "provider_excel_remote_apply_service."
    "requests.put",
    side_effect=persistent_responses,
) as put_mock:
    auth_mock.return_value.get_access_token.return_value = "test-token"
    with patch(
        "app.services."
        "provider_excel_remote_apply_service."
        "time.sleep"
    ):
        try:
            retry_service.upload_content(
                metadata={
                    "drive_id": "DRIVE-1",
                    "item_id": "ITEM-1",
                },
                content=b"xlsx-content",
            )

        except ProviderExcelRemoteLockedError as exc:
            persistent_lock_ok = (
                "bloqueado"
                in str(exc).lower()
                and put_mock.call_count
                == retry_service.UPLOAD_MAX_RETRIES
            )

if persistent_lock_ok:
    ok(
        "HTTP 423 persistente aborta "
        "con error especifico y seguro"
    )
else:
    fail(
        "HTTP 423 persistente no fue "
        "manejado correctamente"
    )


# ------------------------------------------------------------------
# COPIA LOCAL - NO ESCRIBE SHAREPOINT
# ------------------------------------------------------------------

with tempfile.TemporaryDirectory() as temp_dir:
    local_status = FakeStatusService(
        original
    )

    local_service = FakeApplyService(
        local_status
    )

    local_result = (
        local_service.export_local_copy(
            shared_url=
                "https://example.invalid/share",
            relation_results=relations,
            output_dir=temp_dir,
        )
    )

    local_files_ok = all(
        Path(
            local_result[key]
        ).is_file()
        for key in [
            "backup_path",
            "candidate_path",
            "report_path",
        ]
    )

    if (
        local_result["ok"]
        and not local_result[
            "remote_write"
        ]
        and local_result[
            "changes_applied"
        ] == 2
        and local_result[
            "created_1"
        ] == 1
        and local_result[
            "created_2"
        ] == 1
        and local_service.upload_calls == 0
        and local_files_ok
    ):
        ok(
            "Copia local no realiza upload remoto "
            "y genera Excel verificado"
        )
    else:
        fail(
            "La copia local no cumplio el contrato "
            "de seguridad esperado"
        )


print()
print("=" * 100)

if errors:
    print(
        f"RESULTADO: ERROR "
        f"({len(errors)} problema(s))"
    )

    for error in errors:
        print(
            f" - {error}"
        )

    raise SystemExit(1)

print("RESULTADO: OK")
print(
    "Aplicacion Fase 2 PER segura y verificable."
)
print("=" * 100)
