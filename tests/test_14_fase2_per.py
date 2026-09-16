from __future__ import annotations

from io import BytesIO
from pathlib import Path
import sys
import tempfile
from unittest.mock import patch

import pandas as pd
from openpyxl import (
    Workbook,
    load_workbook,
)


ROOT = Path(__file__).resolve().parents[1]

if str(ROOT) not in sys.path:
    sys.path.insert(
        0,
        str(ROOT),
    )


from app.services.provider_creation_status_service import (
    ProviderCreationStatusService,
)
from app.services.provider_phase2_service import (
    ProviderPhase2Service,
)
from app.services.provider_excel_remote_status_service import (
    ProviderExcelRemoteStatusService,
)


errors = []


def ok(message):
    print(
        f"[OK] {message}"
    )


def fail(message):
    print(
        f"[ERROR] {message}"
    )
    errors.append(
        message
    )


def build_remote_bytes():
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Proveedores"

    worksheet.append([
        "nombre",
        "apellido",
        "apellidomaterno",
        "nombreCliente",
        "rucProveedor",
        "nombreProveedor",
        "correo",
        "tipoUsuario",
        "docIdentidad",
        "creado",
    ])

    worksheet.append([
        "ANA",
        "PEREZ",
        "",
        "DOLLARCITY",
        "20111111111",
        "PROVEEDOR A",
        "a@test.com",
        "PROVEEDOR",
        "12345678",
        "",
    ])

    worksheet.append([
        "ANA",
        "PEREZ",
        "",
        "DOLLARCITY",
        "20222222222",
        "PROVEEDOR A2",
        "a@test.com",
        "PROVEEDOR",
        "12345678",
        0,
    ])

    worksheet.append([
        "BETA",
        "LOPEZ",
        "",
        "DOLLARCITY",
        "20333333333",
        "PROVEEDOR B",
        "b@test.com",
        "PROVEEDOR",
        "87654321",
        "",
    ])

    worksheet.append([
        "CARLOS",
        "ERROR",
        "",
        "DOLLARCITY",
        "20444444444",
        "",
        "c@test.com",
        "PROVEEDOR",
        "11112222",
        0,
    ])

    output = BytesIO()
    workbook.save(
        output
    )
    workbook.close()

    return output.getvalue()


def build_formula_workbooks():
    formula_workbook = Workbook()
    formula_sheet = formula_workbook.active
    formula_sheet.title = "Proveedores"

    value_workbook = Workbook()
    value_sheet = value_workbook.active
    value_sheet.title = "Proveedores"

    headers = [
        "nombre",
        "apellido",
        "apellidomaterno",
        "nombreCliente",
        "rucProveedor",
        "nombreProveedor",
        "correo",
        "tipoUsuario",
        "docIdentidad",
        "creado",
    ]

    formula_sheet.append(headers)
    value_sheet.append(headers)

    formula_sheet.append([
        "ANA",
        "PEREZ",
        "",
        "=+D1",
        "20111111111",
        "PROVEEDOR A",
        "a@test.com",
        "=+H1",
        "12345678",
        "",
    ])

    value_sheet.append([
        "ANA",
        "PEREZ",
        "",
        "DOLLARCITY",
        "20111111111",
        "PROVEEDOR A",
        "a@test.com",
        "PROVEEDOR",
        "12345678",
        "",
    ])

    return (
        formula_workbook,
        value_workbook,
    )


print("=" * 100)
print(
    "TEST 14 - FASE 2 PER "
    "Y PREVIEW SHAREPOINT"
)
print("=" * 100)


status_service = (
    ProviderCreationStatusService()
)

if (
    status_service.normalize("")
    == 0
    and status_service.normalize(None)
    == 0
):
    ok(
        "creado vacio se interpreta "
        "como pendiente"
    )
else:
    fail(
        "creado vacio no se interpreta "
        "como pendiente"
    )


with tempfile.TemporaryDirectory() as temp_dir:
    folder = Path(
        temp_dir
    )

    manifest = pd.DataFrame([
        {
            "_item_id":
                "EXCEL-000001",
            "email":
                "a@test.com",
        },
        {
            "_item_id":
                "EXCEL-000003",
            "email":
                "b@test.com",
        },
    ])

    manifest.to_excel(
        folder
        / "usuarios_enviados.xlsx",
        index=False,
    )

    normalized = pd.DataFrame([
        {
            "_item_id":
                "EXCEL-000001",
            "email":
                "a@test.com",
            "estado":
                "OK",
            "creado":
                0,
            "nombre":
                "ANA",
            "apellido_pat":
                "PEREZ",
            "apellido_mat":
                "",
            "nombre_proveedor":
                "PROVEEDOR A",
            "nombre_cliente":
                "DOLLARCITY",
            "_origen_nombre":
                "ANA",
            "_origen_apellido":
                "PEREZ",
            "_origen_apellidomaterno":
                "",
            "_origen_nombreCliente":
                "DOLLARCITY",
            "_origen_rucProveedor":
                "20111111111",
            "_origen_nombreProveedor":
                "PROVEEDOR A",
            "_origen_correo":
                "a@test.com",
            "_origen_tipoUsuario":
                "PROVEEDOR",
            "_origen_docIdentidad":
                "12345678",
        },
        {
            "_item_id":
                "EXCEL-000002",
            "email":
                "a@test.com",
            "estado":
                "OK",
            "creado":
                0,
            "nombre":
                "ANA",
            "apellido_pat":
                "PEREZ",
            "apellido_mat":
                "",
            "nombre_proveedor":
                "PROVEEDOR A2",
            "nombre_cliente":
                "DOLLARCITY",
            "_origen_nombre":
                "ANA",
            "_origen_apellido":
                "PEREZ",
            "_origen_apellidomaterno":
                "",
            "_origen_nombreCliente":
                "DOLLARCITY",
            "_origen_rucProveedor":
                "20222222222",
            "_origen_nombreProveedor":
                "PROVEEDOR A2",
            "_origen_correo":
                "a@test.com",
            "_origen_tipoUsuario":
                "PROVEEDOR",
            "_origen_docIdentidad":
                "12345678",
        },
        {
            "_item_id":
                "EXCEL-000003",
            "email":
                "b@test.com",
            "estado":
                "OK",
            "creado":
                0,
            "nombre":
                "BETA",
            "apellido_pat":
                "LOPEZ",
            "apellido_mat":
                "",
            "nombre_proveedor":
                "PROVEEDOR B",
            "nombre_cliente":
                "DOLLARCITY",
            "_origen_nombre":
                "BETA",
            "_origen_apellido":
                "LOPEZ",
            "_origen_apellidomaterno":
                "",
            "_origen_nombreCliente":
                "DOLLARCITY",
            "_origen_rucProveedor":
                "20333333333",
            "_origen_nombreProveedor":
                "PROVEEDOR B",
            "_origen_correo":
                "b@test.com",
            "_origen_tipoUsuario":
                "PROVEEDOR",
            "_origen_docIdentidad":
                "87654321",
        },
        {
            "_item_id":
                "EXCEL-000004",
            "email":
                "c@test.com",
            "estado":
                "ERROR",
            "creado":
                2,
            "nombre":
                "CARLOS",
            "apellido_pat":
                "ERROR",
            "apellido_mat":
                "",
            "nombre_proveedor":
                "",
            "nombre_cliente":
                "DOLLARCITY",
            "observaciones":
                "NOMBRE PROVEEDOR OBLIGATORIO",
            "_origen_nombre":
                "CARLOS",
            "_origen_apellido":
                "ERROR",
            "_origen_apellidomaterno":
                "",
            "_origen_nombreCliente":
                "DOLLARCITY",
            "_origen_rucProveedor":
                "20444444444",
            "_origen_nombreProveedor":
                "",
            "_origen_correo":
                "c@test.com",
            "_origen_tipoUsuario":
                "PROVEEDOR",
            "_origen_docIdentidad":
                "11112222",
        },
    ])

    normalized.to_excel(
        folder
        / "excel_proveedores_normalizado.xlsx",
        index=False,
    )

    portal = pd.DataFrame([
        {
            "Email":
                "b@test.com",
            "Mensaje de Error":
                (
                    "El usuario ya existe "
                    "en la base de datos."
                ),
        },
    ])

    portal_path = (
        folder
        / "resultado_portal.xlsx"
    )

    portal.to_excel(
        portal_path,
        index=False,
    )

    phase2 = (
        ProviderPhase2Service()
        .preview(
            run_folder=
                str(folder),
            portal_error_excel=
                str(portal_path),
        )
    )

    summary = phase2[
        "summary"
    ]

    if (
        summary[
            "cuentas_enviadas"
        ] == 2
        and summary[
            "cuentas_disponibles"
        ] == 2
        and summary[
            "cuentas_no_creadas"
        ] == 0
    ):
        ok(
            "Portal interpreta creado "
            "y usuario existente como disponible"
        )
    else:
        fail(
            "Resultado de cuentas Portal incorrecto"
        )

    if (
        summary[
            "relaciones_creado_1"
        ] == 3
        and summary[
            "relaciones_creado_2"
        ] == 1
        and summary[
            "relaciones_sin_resultado"
        ] == 0
    ):
        ok(
            "Resultado de cuenta se propaga "
            "a todas las relaciones validas"
        )
    else:
        fail(
            "Propagacion de relaciones incorrecta"
        )

    relation_results = (
        phase2[
            "relation_results"
        ]
    )

    duplicated = (
        relation_results[
            relation_results[
                "email"
            ].eq(
                "a@test.com"
            )
        ]
    )

    if (
        len(duplicated) == 2
        and (
            pd.to_numeric(
                duplicated[
                    "creado_objetivo"
                ],
                errors="coerce",
            )
            .eq(1)
            .all()
        )
    ):
        ok(
            "Correo repetido habilita "
            "todas sus relaciones validas"
        )
    else:
        fail(
            "Correo repetido no propaga creado=1"
        )

    remote_service = (
        ProviderExcelRemoteStatusService()
    )

    remote = (
        remote_service
        .build_preview(
            content=
                build_remote_bytes(),
            relation_results=
                relation_results,
        )
    )

    if (
        remote["ready"]
        and remote[
            "summary"
        ][
            "actualizaciones"
        ] == 4
        and remote[
            "summary"
        ][
            "errores_verificacion"
        ] == 0
        and remote[
            "summary"
        ][
            "conflictos"
        ] == 0
    ):
        ok(
            "Preview remoto verifica "
            "fila antes de cambiar creado"
        )
    else:
        fail(
            "Preview remoto no quedo seguro"
        )

    preview_workbook = (
        load_workbook(
            BytesIO(
                remote[
                    "preview_content"
                ]
            )
        )
    )

    preview_sheet = (
        preview_workbook[
            "Proveedores"
        ]
    )

    statuses = [
        preview_sheet.cell(
            row=row_number,
            column=10,
        ).value
        for row_number
        in range(
            2,
            6,
        )
    ]

    preview_workbook.close()

    if statuses == [
        1,
        1,
        1,
        2,
    ]:
        ok(
            "Copia PREVIEW contiene "
            "los estados esperados"
        )
    else:
        fail(
            "Estados de la copia PREVIEW incorrectos"
        )

    formula_workbook, value_workbook = (
        build_formula_workbooks()
    )

    formula_relation = (
        relation_results[
            relation_results[
                "_item_id"
            ].eq(
                "EXCEL-000001"
            )
        ]
        .copy()
    )

    with patch(
        "app.services.provider_excel_remote_status_service.load_workbook",
        side_effect=[
            formula_workbook,
            value_workbook,
        ],
    ):
        formula_preview = (
            remote_service
            .build_preview(
                content=b"xlsx",
                relation_results=
                    formula_relation,
            )
        )

    formula_output = load_workbook(
        BytesIO(
            formula_preview[
                "preview_content"
            ]
        ),
        data_only=False,
    )

    formula_sheet = formula_output[
        "Proveedores"
    ]

    formula_ok = (
        formula_preview["ready"]
        and formula_preview[
            "summary"
        ][
            "errores_verificacion"
        ] == 0
        and formula_sheet["D2"].value
        == "=+D1"
        and formula_sheet["H2"].value
        == "=+H1"
        and formula_sheet["J2"].value
        == 1
    )

    formula_output.close()

    if formula_ok:
        ok(
            "Verificacion usa valor calculado "
            "y conserva formulas"
        )
    else:
        fail(
            "Formula remota no se verifico "
            "o no se conservo"
        )

    conflict_bytes = (
        build_remote_bytes()
    )

    conflict_workbook = (
        load_workbook(
            BytesIO(
                conflict_bytes
            )
        )
    )

    conflict_sheet = (
        conflict_workbook[
            "Proveedores"
        ]
    )

    conflict_sheet[
        "J2"
    ] = 2

    conflict_output = BytesIO()

    conflict_workbook.save(
        conflict_output
    )

    conflict_workbook.close()

    conflict = (
        remote_service
        .build_preview(
            content=
                conflict_output.getvalue(),
            relation_results=
                relation_results,
        )
    )

    if (
        not conflict["ready"]
        and conflict[
            "summary"
        ][
            "conflictos"
        ] == 1
    ):
        ok(
            "Estado remoto cerrado "
            "no se sobrescribe"
        )
    else:
        fail(
            "Conflicto remoto no fue bloqueado"
        )


print()
print("=" * 100)

if errors:
    print(
        f"RESULTADO: ERROR "
        f"({len(errors)} problema(s))"
    )

    for error in errors:
        print(
            f" - {error}"
        )

    raise SystemExit(1)


print("RESULTADO: OK")
print(
    "FASE 2 PER y PREVIEW "
    "SharePoint validados."
)
print("=" * 100)
