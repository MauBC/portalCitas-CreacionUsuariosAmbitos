from _bootstrap import ROOT

import os

import pandas as pd
from openpyxl import load_workbook

from app.config.portal_template_config import COUNTRY_TEMPLATE_CONFIG
from app.phase3_pipeline import run_phase3_preview
from app.services.client_assignment_service import ClientAssignmentService
from app.services.db_validation_service import DbValidationService


print("=" * 110)
print("TEST 03 - PIPELINE INTEGRADO DE AMBITOS")
print("=" * 110)

errors = []
warnings = []


def ok(message):
    print(f"[OK] {message}")


def fail(message):
    print(f"[ERROR] {message}")
    errors.append(message)


def warn(message):
    print(f"[WARN] {message}")
    warnings.append(message)


def section(title):
    print()
    print("-" * 110)
    print(title)
    print("-" * 110)


def normalize(value):
    return " ".join(str(value or "").strip().upper().split())


def digits(value):
    return "".join(ch for ch in str(value or "") if ch.isdigit())


def parse_groups(value):
    if value is None:
        return set()

    try:
        if pd.isna(value):
            return set()
    except (TypeError, ValueError):
        pass

    result = set()

    for item in str(value).split(","):
        item = item.strip()

        if item.isdigit():
            result.add(int(item))

    return result


# 1. Resolucion de clientes
section("1. RESOLUCION DE CLIENTES")

resolver = ClientAssignmentService()
country_config = COUNTRY_TEMPLATE_CONFIG["SLV"]

cases = [
    ("DoctorSV", "DOCTOR SV"),
    ("DR SV", "DOCTOR SV"),
    ("Telemedicina.", "DOCTOR SV"),
    ("CALLEJA", "CALLEJA S.A"),
    ("CALLEJAS", "CALLEJA S.A"),
    ("Calleja S.A.", "CALLEJA S.A"),
]

for entered, expected in cases:
    detail = resolver.resolve_with_details(
        row={
            "pais": "SLV",
            "nombre_cliente": entered,
        },
        country_config=country_config,
    )

    if (
        detail["status"] == "OK"
        and detail["resolved_client"] == expected
    ):
        ok(f"{entered!r} -> {expected}")
    else:
        fail(
            f"{entered!r}: esperado {expected}, "
            f"obtenido {detail}"
        )

detail = resolver.resolve_with_details(
    row={
        "pais": "SLV",
        "nombre_cliente": "cliente completamente desconocido xyz",
    },
    country_config=country_config,
)

if detail["status"] == "REVISAR":
    ok("Cliente desconocido queda para revision manual")
else:
    fail("Cliente desconocido fue asignado automaticamente")


# 2. Clientes en PostgreSQL
section("2. CLIENTES EN POSTGRESQL")

db_clients = []

try:
    db = DbValidationService()
    db_clients = db.get_persons_by_names(
        ["DOCTOR SV", "CALLEJA S.A"]
    )

    by_name = {
        normalize(row.get("name")): row
        for row in db_clients
    }

    for client_name in ["CALLEJA S.A", "DOCTOR SV"]:
        row = by_name.get(client_name)

        if not row:
            fail(f"{client_name} no fue encontrado en BD")
            continue

        if not digits(row.get("document_number")):
            fail(f"{client_name} no tiene document_number")
            continue

        ok(f"{client_name} existe en BD con identificador")

except Exception as exc:
    fail("Consulta BD fallo: " + str(exc))


# 3. Ejecutar FASE 3
section("3. GENERACION DE AMBITOS")

result = run_phase3_preview()

for label, key in [
    ("Pais", "country"),
    ("Registros fuente", "usuarios_fuente"),
    ("CREADO=1 procesados", "usuarios"),
    ("Filtrados CREADO!=1", "filtrados_creado"),
    ("Matched", "matched"),
    ("Revision manual", "no_match"),
    ("Entidades", "entidades"),
    ("RelacionNueva", "relacion_nueva"),
    ("Ambitos", "ambitos"),
]:
    print(f"{label:<22}: {result.get(key)}")

usuarios = int(result.get("usuarios", 0))
matched = int(result.get("matched", 0))
no_match = int(result.get("no_match", 0))

if usuarios <= 0:
    fail("No hay registros CREADO=1 para FASE 3")
elif matched + no_match == usuarios:
    ok("Matched + revision manual = registros procesados")
else:
    fail("Resultado FASE 3 inconsistente")

if no_match:
    warn(
        f"{no_match} relacion(es) requieren "
        "asignar cliente manualmente"
    )
else:
    ok("No existen clientes pendientes de revision manual")


# 4. Archivos
section("4. ARCHIVOS DE AMBITOS")

template_path = result.get("template_path")
diagnostic_path = result.get("diagnostics_path")
review_path = result.get("review_path")

for name, path in [
    ("PLANTILLA", template_path),
    ("DIAGNOSTICO", diagnostic_path),
]:
    if path and os.path.exists(path):
        ok(f"{name}: {path}")
    else:
        fail(f"{name} no fue generado")

if no_match:
    if review_path and os.path.exists(review_path):
        ok("Excel de revision manual generado")
    else:
        fail("Hay revision manual pero falta su Excel")
elif review_path:
    fail("Se genero revision manual sin casos pendientes")
else:
    ok("No se genero Excel manual porque no era necesario")


# 5. Cargar plantilla
section("5. ESTRUCTURA PLANTILLA AMBITOS")

expected_sheets = [
    "Entidades",
    "Ambitos",
    "RelacionNueva",
    "TipoNegocioNuevo",
    "TipoNegocioEliminar",
    "RelacionEliminar",
    "SedeNueva",
    "SedeEliminar",
]

excel = pd.ExcelFile(template_path)

if excel.sheet_names == expected_sheets:
    ok("Orden y hojas de plantilla oficial correctos")
else:
    fail("Hojas u orden incorrectos: " + str(excel.sheet_names))

entidades = pd.read_excel(
    template_path,
    sheet_name="Entidades",
    dtype=str,
)

ambitos = pd.read_excel(
    template_path,
    sheet_name="Ambitos",
    dtype=str,
)

relaciones = pd.read_excel(
    template_path,
    sheet_name="RelacionNueva",
    dtype=str,
)

tipo_negocio = pd.read_excel(
    template_path,
    sheet_name="TipoNegocioNuevo",
    dtype=str,
)

sedes = pd.read_excel(
    template_path,
    sheet_name="SedeNueva",
    dtype=str,
)

for name, frame, key in [
    ("Entidades", entidades, "entidades"),
    ("Ambitos", ambitos, "ambitos"),
    ("RelacionNueva", relaciones, "relacion_nueva"),
]:
    if len(frame) == int(result.get(key, 0)):
        ok(f"Cantidad {name} correcta")
    else:
        fail(f"Cantidad {name} incorrecta")


# 6. TipoNegocioNuevo
section("6. TIPO NEGOCIO NUEVO")

expected_rubros = {
    1: "TECNOLOGÍA",
    2: "CONSUMO MASIVO",
    3: "PERFUMERÍA Y FARMACIA",
    4: "RETAIL",
    5: "ALIMENTOS Y BEBIDAS",
    6: "PRODUCTOS ADHESIVOS",
    7: "TRATAMIENTO Y DISTRIBUCIÓN DE AGUA",
    8: "FARMACIA Y COSMÉTICOS",
    9: "MINERÍA",
}

actual = {}

for _, row in tipo_negocio.iterrows():
    try:
        group = int(float(row.get("Grupo")))
    except (TypeError, ValueError):
        continue

    actual[group] = (
        normalize(row.get("Tipo de negocio")),
        normalize(row.get("Rubro")),
    )

if set(actual) != set(range(1, 10)):
    fail("TipoNegocioNuevo no contiene exactamente grupos 1 al 9")
else:
    for group, rubro in expected_rubros.items():
        tipo, actual_rubro = actual[group]

        if tipo != "CENTROS DE DISTRIBUCIÓN":
            fail(f"Grupo {group}: Tipo de negocio incorrecto")

        if actual_rubro != rubro:
            fail(
                f"Grupo {group}: esperado {rubro!r}, "
                f"obtenido {actual_rubro!r}"
            )

    if not errors:
        ok("Catalogo de 9 grupos correcto")


# 7. SedeNueva
section("7. SEDES SLV")

actual_sedes = {}

for _, row in sedes.iterrows():
    try:
        group = int(float(row.get("Grupo")))
    except (TypeError, ValueError):
        continue

    actual_sedes[group] = (
        normalize(row.get("Sede")),
        normalize(row.get("Tipo de negocio")),
    )

expected_sedes = {
    1: ("RANSA APOPA", "FRÍOS"),
    2: ("RANSA SANTA ELENA", "SECOS"),
}

if actual_sedes == expected_sedes:
    ok("1=RANSA APOPA/FRIOS y 2=RANSA SANTA ELENA/SECOS")
else:
    fail("SedeNueva incorrecta: " + str(actual_sedes))


# 8. RelacionNueva y Entidades
section("8. RELACION NUEVA Y ENTIDADES")

relation_groups = set()

for value in relaciones.get("Grupo", pd.Series(dtype=str)).dropna():
    try:
        relation_groups.add(int(float(value)))
    except ValueError:
        pass

if relation_groups.issubset({1, 2}):
    ok("RelacionNueva usa solo grupos 1 y 2")
else:
    fail("RelacionNueva contiene grupos no permitidos")

if not relaciones.empty:
    duplicate_relation = (
        relaciones.assign(
            _grupo=relaciones["Grupo"].astype(str),
            _ruc=relaciones["RUC"].fillna("").apply(digits),
        )
        .duplicated(subset=["_grupo", "_ruc"])
        .any()
    )

    if duplicate_relation:
        fail("RelacionNueva duplica Grupo + RUC")
    else:
        ok("RelacionNueva no duplica Grupo + RUC")

provider_rucs = {
    digits(value)
    for value in relaciones.get("RUC", pd.Series(dtype=str)).dropna()
    if digits(value)
}

provider_entities = entidades[
    entidades["RUC"]
    .fillna("")
    .apply(digits)
    .isin(provider_rucs)
].copy()

if (
    provider_entities["RUC"]
    .fillna("")
    .apply(digits)
    .duplicated()
    .any()
):
    fail("Entidades duplica empresa por NIT/RUC")
else:
    ok("Una sola empresa en Entidades por NIT/RUC")

for _, row in provider_entities.iterrows():
    name = str(row.get("Razon Social") or "").strip()

    if name != name.upper():
        fail(f"Empresa no esta en mayusculas: {name}")

for _, relation in relaciones.iterrows():
    ruc = digits(relation.get("RUC"))

    try:
        group = int(float(relation.get("Grupo")))
    except (TypeError, ValueError):
        continue

    entity_rows = provider_entities[
        provider_entities["RUC"]
        .fillna("")
        .apply(digits)
        .eq(ruc)
    ]

    if entity_rows.empty:
        fail(f"No existe Entidad para proveedor {ruc}")
        continue

    business_groups = parse_groups(
        entity_rows.iloc[0].get("Tipo de negocio Nuevo")
    )

    expected_group = 5 if group == 1 else 3

    if expected_group not in business_groups:
        fail(
            f"Proveedor {ruc}: grupo cliente {group}, "
            f"falta TipoNegocioNuevo {expected_group}"
        )


# 9. Clientes en Entidades
section("9. CLIENTES EN ENTIDADES")

db_by_name = {
    normalize(row.get("name")): row
    for row in db_clients
}

for client_name, expected_group in {
    "CALLEJA S.A": 1,
    "DOCTOR SV": 2,
}.items():
    db_row = db_by_name.get(client_name)

    if not db_row:
        continue

    client_ruc = digits(db_row.get("document_number"))

    rows = entidades[
        entidades["RUC"]
        .fillna("")
        .apply(digits)
        .eq(client_ruc)
    ]

    if rows.empty:
        if expected_group in relation_groups:
            fail(f"Falta {client_name} en Entidades")
        continue

    try:
        actual_group = int(float(rows.iloc[0].get("Relación Nueva")))
    except (TypeError, ValueError):
        actual_group = None

    if actual_group == expected_group:
        ok(f"{client_name}: Relacion Nueva = {expected_group}")
    else:
        fail(
            f"{client_name}: esperado {expected_group}, "
            f"obtenido {actual_group}"
        )



# =============================================================================
# TIPO DOCUMENTO POR PAIS
# =============================================================================

expected_entity_type = {
    "SLV": "NIT",
    "PER": "RUC",
}.get(
    str(
        result.get("country")
        or ""
    ).strip().upper()
)

if not expected_entity_type:
    fail(
        "Pais sin Tipo de documento "
        "configurado para Entidades"
    )

elif entidades.empty:
    fail(
        "No existen Entidades para "
        "validar Tipo de documento"
    )

else:
    actual_types = {
        str(value or "")
        .strip()
        .upper()
        for value in (
            entidades["Tipo"]
            .fillna("")
        )
        if str(value or "").strip()
    }

    if actual_types == {
        expected_entity_type
    }:
        ok(
            "Entidades usa Tipo="
            f"{expected_entity_type}"
        )
    else:
        fail(
            "Tipo de documento incorrecto: "
            f"{sorted(actual_types)}"
        )

    if "Tipo Usuario" not in entidades.columns:
        fail(
            "Entidades no contiene "
            "la columna Tipo Usuario"
        )
    else:
        provider_rucs_tipo = set(
            relaciones["RUC"]
            .fillna("")
            .astype(str)
            .str.replace(
                r"\D",
                "",
                regex=True,
            )
            .loc[
                lambda series:
                    series.ne("")
            ]
            .tolist()
        )

        entity_rucs_tipo = (
            entidades["RUC"]
            .fillna("")
            .astype(str)
            .str.replace(
                r"\D",
                "",
                regex=True,
            )
        )

        provider_mask_tipo = (
            entity_rucs_tipo
            .isin(
                provider_rucs_tipo
            )
        )

        provider_rows_tipo = (
            entidades[
                provider_mask_tipo
            ]
        )

        client_rows_tipo = (
            entidades[
                ~provider_mask_tipo
            ]
        )

        provider_types = (
            provider_rows_tipo[
                "Tipo Usuario"
            ]
            .fillna("")
            .astype(str)
            .str.strip()
            .str.upper()
        )

        client_types = (
            client_rows_tipo[
                "Tipo Usuario"
            ]
            .fillna("")
            .astype(str)
            .str.strip()
            .str.upper()
        )

        if provider_rows_tipo.empty:
            fail(
                "No existen entidades proveedoras "
                "para validar Tipo Usuario"
            )
        elif provider_types.eq(
            "PROVEEDOR"
        ).all():
            ok(
                "Entidades proveedoras usan "
                "Tipo Usuario=PROVEEDOR"
            )
        else:
            invalid_values = sorted(
                set(
                    provider_types[
                        provider_types.ne(
                            "PROVEEDOR"
                        )
                    ]
                )
            )

            fail(
                "Tipo Usuario incorrecto "
                "en proveedores: "
                + str(
                    invalid_values
                )
            )

        if client_rows_tipo.empty:
            fail(
                "No existen entidades cliente "
                "para validar Tipo Usuario"
            )
        elif client_types.eq("CLIENTE").all():
            ok(
                "Entidades cliente conservan "
                "Tipo Usuario=CLIENTE"
            )
        else:
            fail(
                "Las entidades cliente deben "
                "tener Tipo Usuario=CLIENTE"
            )

# 10. Ambitos
section("10. AMBITOS")

if ambitos.empty:
    fail("No se generaron ambitos")
else:
    duplicate_ambito = (
        ambitos.assign(
            _email=ambitos["Email"].fillna("").str.strip().str.lower(),
            _empresa=ambitos["Empresa"].fillna("").apply(digits),
        )
        .duplicated(subset=["_email", "_empresa"])
        .any()
    )

    if duplicate_ambito:
        fail("Ambitos duplica Email + Empresa")
    else:
        ok("Ambitos tiene una fila por Email + Empresa")

    invalid = []

    for _, row in ambitos.iterrows():
        groups = parse_groups(row.get("Sedes Nuevas"))

        if not groups or not groups.issubset({1, 2}):
            invalid.append(
                (
                    row.get("Email"),
                    row.get("Empresa"),
                    row.get("Sedes Nuevas"),
                )
            )

    if invalid:
        fail("Sedes Nuevas invalidas: " + str(invalid[:5]))
    else:
        ok("Ambitos usa Sede Nueva 1, 2 o 1,2")


# 11. Revision manual
section("11. REVISION MANUAL")

if no_match:
    workbook = load_workbook(review_path)
    worksheet = workbook["REVISION_CLIENTES"]

    headers = [
        worksheet.cell(row=1, column=index).value
        for index in range(1, 11)
    ]

    required = {
        "CLIENTE RECIBIDO",
        "CLIENTE ASIGNADO",
        "MOTIVO",
    }

    if required.issubset(set(headers)):
        ok("Excel manual conserva datos de revision")
    else:
        fail("Excel manual tiene columnas incorrectas")

    if worksheet.data_validations.dataValidation:
        ok("CLIENTE ASIGNADO tiene desplegable")
    else:
        fail("Falta desplegable de cliente")

    if worksheet["H2"].fill.fill_type == "solid":
        ok("CLIENTE ASIGNADO esta resaltado")
    else:
        fail("CLIENTE ASIGNADO no esta resaltado")


print()
print("=" * 110)

if warnings:
    print(f"WARNINGS: {len(warnings)}")

if errors:
    print(f"RESULTADO: ERROR ({len(errors)} problema(s))")

    for error in errors:
        print(f" - {error}")

    raise SystemExit(1)

print("RESULTADO: OK")
print("Pipeline de ambitos funcionando con reglas SLV.")
print("=" * 110)
