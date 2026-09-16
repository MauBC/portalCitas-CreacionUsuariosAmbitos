"""Offline smoke check usable both from source and the packaged executable."""
import json
from pathlib import Path
import traceback


def run_self_check(report_path: str) -> int:
    report = {"ok": False, "checks": []}
    try:
        from PySide6.QtWidgets import QApplication
        from openpyxl import load_workbook
        from app.config.paths import CONFIG_ROOT, PROJECT_ROOT, TEMPLATES_ROOT, get_runtime_root
        from app.ui.main_window import MainWindow
        from app.ui.styles import APP_STYLE
        # Import every pipeline to detect missing bundled dependencies. No client
        # is instantiated and no pipeline is executed by this diagnostic.
        from app import (main_pipeline, provider_excel_pipeline, phase2_pipeline,
                         provider_phase2_pipeline, provider_phase2_apply_pipeline, phase3_pipeline)

        for name in ("subida_usuarios.xlsx", "subida_ambitos.xlsx"):
            workbook = load_workbook(TEMPLATES_ROOT / name, read_only=True)
            assert workbook.sheetnames, f"Plantilla vacia: {name}"
            workbook.close()
            report["checks"].append(name)
        with (CONFIG_ROOT / "ambitos_config.json").open(encoding="utf-8-sig") as stream:
            assert json.load(stream), "Configuracion de ambitos vacia"
        assert (PROJECT_ROOT / "Ransalogo.ico").is_file(), "Falta icono"
        report["checks"].append("configuracion e icono")
        app = QApplication.instance() or QApplication([])
        app.setStyleSheet(APP_STYLE)
        window = MainWindow()
        try:
            window.show()
            for size in ((1020, 680), (1400, 900)):
                window.resize(*size)
                for index in range(window.country_combo.count()):
                    window.country_combo.setCurrentIndex(index)
                    for page in range(window.pages.count()):
                        window.pages.setCurrentIndex(page)
                        app.processEvents()
            report["checks"].append("GUI: paises, paginas y resize")
        finally:
            window.close()
        report.update(ok=True, runtime_root=str(get_runtime_root()))
    except Exception:
        from app.services.error_report_service import redact_secrets
        report["error"] = redact_secrets(traceback.format_exc())
    destination = Path(report_path).resolve()
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
    return 0 if report["ok"] else 1
