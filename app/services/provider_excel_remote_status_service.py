from __future__ import annotations

from base64 import urlsafe_b64encode
from io import BytesIO
from pathlib import Path
import re

import pandas as pd
import requests
from openpyxl import load_workbook

from app.auth.msal_auth import (
    MsalAuthService,
)


class ProviderExcelRemoteStatusService:
    SHEET_NAME = "Proveedores"

    VERIFY_FIELDS = [
        (
            "_origen_nombre",
            "nombre",
        ),
        (
            "_origen_apellido",
            "apellido",
        ),
        (
            "_origen_apellidomaterno",
            "apellidomaterno",
        ),
        (
            "_origen_nombreCliente",
            "nombreCliente",
        ),
        (
            "_origen_rucProveedor",
            "rucProveedor",
        ),
        (
            "_origen_nombreProveedor",
            "nombreProveedor",
        ),
        (
            "_origen_correo",
            "correo",
        ),
        (
            "_origen_tipoUsuario",
            "tipoUsuario",
        ),
        (
            "_origen_docIdentidad",
            "docIdentidad",
        ),
    ]

    def download(
        self,
        shared_url: str,
    ) -> bytes:
        url = str(
            shared_url or ""
        ).strip()

        if not url:
            raise ValueError(
                "Falta URL compartida "
                "de SharePoint/OneDrive."
            )

        encoded = (
            self._encode_shared_url(
                url
            )
        )

        token = (
            MsalAuthService()
            .get_access_token()
        )

        endpoint = (
            "https://graph.microsoft.com/v1.0/"
            f"shares/{encoded}/driveItem/content"
        )

        with requests.get(
            endpoint,
            headers={
                "Authorization":
                    f"Bearer {token}",
            },
            timeout=120,
            allow_redirects=True,
        ) as response:
            if not response.ok:
                detail = (
                    response.text[:1000]
                    if response.text
                    else ""
                )

                raise RuntimeError(
                    "No se pudo descargar "
                    "el Excel remoto. "
                    f"HTTP {response.status_code}. "
                    f"{detail}"
                )

            content = response.content

        if not content:
            raise RuntimeError(
                "Graph devolvio un "
                "archivo remoto vacio."
            )

        return content

    def build_preview(
        self,
        content: bytes,
        relation_results:
            pd.DataFrame,
    ) -> dict:
        if not content:
            raise ValueError(
                "El contenido del Excel "
                "remoto esta vacio."
            )

        required_results = {
            "_item_id",
            "creado_objetivo",
        }

        missing_results = (
            required_results
            - set(
                relation_results.columns
            )
        )

        if missing_results:
            raise ValueError(
                "Faltan columnas en "
                "resultado Fase 2: "
                f"{sorted(missing_results)}"
            )

        workbook = load_workbook(
            BytesIO(content),
            data_only=False,
        )

        values_workbook = load_workbook(
            BytesIO(content),
            data_only=True,
        )

        try:
            if self.SHEET_NAME not in (
                workbook.sheetnames
            ):
                raise ValueError(
                    "El Excel remoto no contiene "
                    f"la hoja '{self.SHEET_NAME}'."
                )

            worksheet = workbook[
                self.SHEET_NAME
            ]

            values_worksheet = (
                values_workbook[
                    self.SHEET_NAME
                ]
            )

            headers = (
                self._build_header_map(
                    worksheet
                )
            )

            creado_key = (
                self._normalize_header(
                    "creado"
                )
            )

            if creado_key not in headers:
                raise ValueError(
                    "El Excel remoto no contiene "
                    "la columna creado."
                )

            for _, header in (
                self.VERIFY_FIELDS
            ):
                normalized = (
                    self._normalize_header(
                        header
                    )
                )

                if normalized not in headers:
                    raise ValueError(
                        "El Excel remoto no contiene "
                        f"la columna {header}."
                    )

            creado_column = (
                headers[
                    creado_key
                ]
            )

            plan = []

            updates = 0
            skipped = 0
            conflicts = 0
            verification_errors = 0
            invalid_states = 0

            for _, row in (
                relation_results.iterrows()
            ):
                item_id = str(
                    row.get("_item_id")
                    or ""
                ).strip()

                desired = (
                    self._normalize_desired(
                        row.get(
                            "creado_objetivo"
                        )
                    )
                )

                excel_data_index = (
                    self._parse_item_id(
                        item_id
                    )
                )

                excel_row = (
                    excel_data_index + 1
                    if excel_data_index
                    is not None
                    else None
                )

                action = ""
                detail = ""

                if desired not in {
                    1,
                    2,
                }:
                    invalid_states += 1
                    action = (
                        "ERROR_OBJETIVO_INVALIDO"
                    )
                    detail = (
                        "creado_objetivo "
                        "debe ser 1 o 2"
                    )

                    plan.append(
                        self._plan_row(
                            row=row,
                            item_id=item_id,
                            excel_row=excel_row,
                            current="",
                            desired=desired,
                            action=action,
                            detail=detail,
                        )
                    )
                    continue

                if excel_row is None:
                    verification_errors += 1
                    action = (
                        "ERROR_ITEM_ID"
                    )
                    detail = (
                        "Formato _item_id invalido"
                    )

                    plan.append(
                        self._plan_row(
                            row=row,
                            item_id=item_id,
                            excel_row="",
                            current="",
                            desired=desired,
                            action=action,
                            detail=detail,
                        )
                    )
                    continue

                if (
                    excel_row < 2
                    or excel_row
                    > worksheet.max_row
                ):
                    verification_errors += 1
                    action = (
                        "ERROR_FILA_NO_EXISTE"
                    )
                    detail = (
                        "La fila calculada no existe "
                        "en el Excel remoto"
                    )

                    plan.append(
                        self._plan_row(
                            row=row,
                            item_id=item_id,
                            excel_row=excel_row,
                            current="",
                            desired=desired,
                            action=action,
                            detail=detail,
                        )
                    )
                    continue

                mismatches = (
                    self._verify_business_row(
                        worksheet=
                            values_worksheet,
                        excel_row=
                            excel_row,
                        headers=
                            headers,
                        relation_row=
                            row,
                    )
                )

                current = (
                    self._normalize_current_status(
                        worksheet.cell(
                            row=excel_row,
                            column=creado_column,
                        ).value
                    )
                )

                if mismatches:
                    verification_errors += 1
                    action = (
                        "ERROR_DATOS_NO_COINCIDEN"
                    )
                    detail = (
                        "Columnas distintas: "
                        + ", ".join(
                            mismatches
                        )
                    )

                elif current is None:
                    invalid_states += 1
                    action = (
                        "ERROR_ESTADO_REMOTO_INVALIDO"
                    )
                    detail = (
                        "creado remoto no es "
                        "vacio, 0, 1 ni 2"
                    )

                elif current == desired:
                    skipped += 1
                    action = (
                        "SIN_CAMBIO"
                    )

                elif current in {
                    1,
                    2,
                }:
                    conflicts += 1
                    action = (
                        "CONFLICTO_CERRADO"
                    )
                    detail = (
                        f"Remoto tiene {current} "
                        f"y Fase 2 requiere {desired}"
                    )

                elif current == 0:
                    worksheet.cell(
                        row=excel_row,
                        column=creado_column,
                    ).value = desired

                    updates += 1
                    action = (
                        "ACTUALIZAR"
                    )

                plan.append(
                    self._plan_row(
                        row=row,
                        item_id=item_id,
                        excel_row=excel_row,
                        current=current,
                        desired=desired,
                        action=action,
                        detail=detail,
                    )
                )

            output = BytesIO()

            workbook.save(
                output
            )

            preview_content = (
                output.getvalue()
            )

        finally:
            workbook.close()
            values_workbook.close()

        ready = (
            verification_errors == 0
            and invalid_states == 0
            and conflicts == 0
        )

        plan_df = pd.DataFrame(
            plan
        )

        summary = {
            "filas_resultado":
                len(
                    relation_results
                ),
            "actualizaciones":
                updates,
            "sin_cambio":
                skipped,
            "conflictos":
                conflicts,
            "errores_verificacion":
                verification_errors,
            "estados_invalidos":
                invalid_states,
            "listo_para_escritura":
                bool(
                    ready
                ),
        }

        return {
            "ready":
                ready,
            "summary":
                summary,
            "plan":
                plan_df,
            "preview_content":
                preview_content,
        }

    def export_preview(
        self,
        shared_url: str,
        relation_results:
            pd.DataFrame,
        output_dir: str,
    ) -> dict:
        folder = Path(
            output_dir
        )

        folder.mkdir(
            parents=True,
            exist_ok=True,
        )

        content = self.download(
            shared_url
        )

        preview = self.build_preview(
            content=content,
            relation_results=
                relation_results,
        )

        preview_path = (
            folder
            / "excel_sharepoint_creado_PREVIEW.xlsx"
        )

        plan_path = (
            folder
            / "fase2_PER_plan_sharepoint.xlsx"
        )

        preview_path.write_bytes(
            preview[
                "preview_content"
            ]
        )

        summary_rows = [
            {
                "METRICA": key,
                "VALOR": value,
            }
            for key, value
            in preview[
                "summary"
            ].items()
        ]

        with pd.ExcelWriter(
            plan_path,
            engine="openpyxl",
        ) as writer:
            pd.DataFrame(
                summary_rows
            ).to_excel(
                writer,
                sheet_name="RESUMEN",
                index=False,
            )

            preview[
                "plan"
            ].to_excel(
                writer,
                sheet_name="PLAN",
                index=False,
            )

        return {
            **preview,
            "preview_path":
                str(
                    preview_path.resolve()
                ),
            "plan_path":
                str(
                    plan_path.resolve()
                ),
        }

    def _verify_business_row(
        self,
        worksheet,
        excel_row: int,
        headers: dict,
        relation_row,
    ) -> list[str]:
        mismatches = []

        for source_column, header in (
            self.VERIFY_FIELDS
        ):
            if (
                source_column
                not in relation_row.index
            ):
                continue

            expected = (
                self._normalize_business_value(
                    relation_row.get(
                        source_column
                    )
                )
            )

            column_index = headers[
                self._normalize_header(
                    header
                )
            ]

            actual = (
                self._normalize_business_value(
                    worksheet.cell(
                        row=excel_row,
                        column=column_index,
                    ).value
                )
            )

            if expected != actual:
                mismatches.append(
                    header
                )

        return mismatches

    @staticmethod
    def _plan_row(
        row,
        item_id,
        excel_row,
        current,
        desired,
        action,
        detail,
    ) -> dict:
        return {
            "_item_id":
                item_id,
            "FILA EXCEL":
                excel_row,
            "CORREO":
                str(
                    row.get("email")
                    or ""
                ).strip().lower(),
            "NOMBRE":
                str(
                    row.get("nombre")
                    or ""
                ).strip(),
            "APELLIDO":
                " ".join(
                    part
                    for part in [
                        str(
                            row.get(
                                "apellido_pat"
                            )
                            or ""
                        ).strip(),
                        str(
                            row.get(
                                "apellido_mat"
                            )
                            or ""
                        ).strip(),
                    ]
                    if part
                ),
            "PROVEEDOR":
                str(
                    row.get(
                        "nombre_proveedor"
                    )
                    or ""
                ).strip(),
            "CLIENTE":
                str(
                    row.get(
                        "nombre_cliente"
                    )
                    or ""
                ).strip(),
            "CREADO ACTUAL":
                current,
            "CREADO NUEVO":
                (
                    desired
                    if desired
                    is not None
                    else ""
                ),
            "ACCION":
                action,
            "DETALLE":
                detail,
        }

    @classmethod
    def _build_header_map(
        cls,
        worksheet,
    ) -> dict:
        result = {}

        for cell in worksheet[1]:
            normalized = (
                cls._normalize_header(
                    cell.value
                )
            )

            if normalized:
                result[
                    normalized
                ] = cell.column

        return result

    @staticmethod
    def _parse_item_id(
        item_id: str,
    ):
        match = re.fullmatch(
            r"EXCEL-(\d{6})",
            str(
                item_id or ""
            ).strip(),
        )

        if not match:
            return None

        value = int(
            match.group(1)
        )

        if value <= 0:
            return None

        return value

    @staticmethod
    def _normalize_desired(
        value,
    ):
        try:
            numeric = float(
                str(value).strip()
            )
        except (
            TypeError,
            ValueError,
        ):
            return None

        if not numeric.is_integer():
            return None

        value = int(
            numeric
        )

        return (
            value
            if value in {
                1,
                2,
            }
            else None
        )

    @staticmethod
    def _normalize_current_status(
        value,
    ):
        if value is None:
            return 0

        try:
            if pd.isna(value):
                return 0
        except (
            TypeError,
            ValueError,
        ):
            pass

        text = str(
            value
        ).strip()

        if not text:
            return 0

        try:
            numeric = float(
                text
            )
        except (
            TypeError,
            ValueError,
        ):
            return None

        if not numeric.is_integer():
            return None

        value = int(
            numeric
        )

        return (
            value
            if value in {
                0,
                1,
                2,
            }
            else None
        )

    @staticmethod
    def _normalize_business_value(
        value,
    ) -> str:
        if value is None:
            return ""

        try:
            if pd.isna(value):
                return ""
        except (
            TypeError,
            ValueError,
        ):
            pass

        if isinstance(
            value,
            float,
        ):
            if value.is_integer():
                value = int(
                    value
                )

        text = str(
            value
        ).strip()

        if (
            text.endswith(".0")
            and text[:-2].isdigit()
        ):
            text = text[:-2]

        return " ".join(
            text.split()
        ).casefold()

    @staticmethod
    def _normalize_header(
        value,
    ) -> str:
        return (
            str(
                value or ""
            )
            .strip()
            .replace(" ", "")
            .replace("_", "")
            .lower()
        )

    @staticmethod
    def _encode_shared_url(
        shared_url: str,
    ) -> str:
        encoded = urlsafe_b64encode(
            shared_url.encode(
                "utf-8"
            )
        ).decode(
            "ascii"
        ).rstrip("=")

        return "u!" + encoded
