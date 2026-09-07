from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import (
    Alignment,
    Font,
    PatternFill,
)
from openpyxl.worksheet.datavalidation import (
    DataValidation,
)

from app.services.ambitos_config_service import (
    AmbitosConfigService,
)


class AmbitosReviewService:

    HEADERS = [
        "ITEM ID",
        "CORREO",
        "PROVEEDOR",
        "NIT",
        "CLIENTE RECIBIDO",
        "CLIENTE SUGERIDO",
        "SCORE",
        "CLIENTE ASIGNADO",
        "ESTADO",
        "MOTIVO",
    ]

    def __init__(self):
        self.config = (
            AmbitosConfigService()
        )

    def export(
        self,
        no_match: list[dict],
        output_path,
        country_code: str = "SLV",
    ):
        if not no_match:
            return None

        output_path = Path(
            output_path
        )

        output_path.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        previous_assignments = (
            self._read_existing_assignments(
                output_path
            )
        )

        workbook = Workbook()

        worksheet = workbook.active
        worksheet.title = (
            "REVISION_CLIENTES"
        )

        catalog = workbook.create_sheet(
            "CATALOGO_CLIENTES"
        )

        client_names = (
            self.config.get_client_names(
                country_code
            )
        )

        catalog["A1"] = "CLIENTE"

        for index, client in enumerate(
            client_names,
            start=2,
        ):
            catalog.cell(
                row=index,
                column=1,
                value=client,
            )

        catalog.sheet_state = "hidden"

        header_fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78",
        )

        header_font = Font(
            color="FFFFFF",
            bold=True,
        )

        review_fill = PatternFill(
            fill_type="solid",
            fgColor="FFF2CC",
        )

        input_fill = PatternFill(
            fill_type="solid",
            fgColor="F4B183",
        )

        for column, header in enumerate(
            self.HEADERS,
            start=1,
        ):
            cell = worksheet.cell(
                row=1,
                column=column,
                value=header,
            )

            cell.fill = header_fill
            cell.font = header_font

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        for excel_row, item in enumerate(
            no_match,
            start=2,
        ):
            item_id = self._normalize_item_id(
                item.get(
                    "_item_id",
                    "",
                )
            )

            original_client = (
                item.get(
                    "cliente_original"
                )
                or item.get(
                    "cliente_input"
                )
                or ""
            )

            assigned = (
                previous_assignments.get(
                    item_id,
                    "",
                )
            )

            status = (
                "MANUAL_ASIGNADO"
                if assigned
                else "REVISION_CLIENTE"
            )

            values = [
                item_id,
                item.get(
                    "email",
                    "",
                ),
                item.get(
                    "proveedor_nombre",
                    "",
                ),
                item.get(
                    "proveedor_ruc",
                    "",
                ),
                original_client,
                item.get(
                    "cliente_sugerido",
                    "",
                ),
                item.get(
                    "match_score",
                    "",
                ),
                assigned,
                status,
                item.get(
                    "mensaje",
                    "",
                ),
            ]

            for column, value in enumerate(
                values,
                start=1,
            ):
                cell = worksheet.cell(
                    row=excel_row,
                    column=column,
                    value=value,
                )

                cell.fill = review_fill

                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

            client_cell = worksheet.cell(
                row=excel_row,
                column=8,
            )

            client_cell.fill = input_fill

            client_cell.font = Font(
                bold=True,
            )

            client_cell.comment = Comment(
                (
                    "Seleccione el cliente "
                    "correcto para esta relacion. "
                    "Si no lo conoce, deje la "
                    "celda vacia."
                ),
                "Sistema",
            )

        if client_names:
            validation = DataValidation(
                type="list",
                formula1=(
                    "'CATALOGO_CLIENTES'!"
                    f"$A$2:$A$"
                    f"{len(client_names) + 1}"
                ),
                allow_blank=True,
            )

            validation.error = (
                "Seleccione un cliente "
                "del catalogo."
            )

            validation.errorTitle = (
                "Cliente invalido"
            )

            worksheet.add_data_validation(
                validation
            )

            validation.add(
                (
                    f"H2:H"
                    f"{len(no_match) + 1}"
                )
            )

        widths = {
            "A": 12,
            "B": 35,
            "C": 40,
            "D": 22,
            "E": 55,
            "F": 25,
            "G": 12,
            "H": 25,
            "I": 22,
            "J": 48,
        }

        for column, width in (
            widths.items()
        ):
            worksheet.column_dimensions[
                column
            ].width = width

        worksheet.freeze_panes = "A2"

        worksheet.auto_filter.ref = (
            f"A1:J{len(no_match) + 1}"
        )

        workbook.save(
            output_path
        )

        return str(
            output_path.resolve()
        )

    def read_assignments(
        self,
        review_path,
        country_code: str,
    ) -> dict:
        review_path = Path(
            review_path
        )

        if not review_path.exists():
            raise FileNotFoundError(
                f"No existe archivo de revision: "
                f"{review_path}"
            )

        workbook = load_workbook(
            review_path,
            data_only=False,
        )

        if (
            "REVISION_CLIENTES"
            not in workbook.sheetnames
        ):
            raise ValueError(
                "El archivo de revision no "
                "contiene REVISION_CLIENTES."
            )

        worksheet = workbook[
            "REVISION_CLIENTES"
        ]

        headers = {
            str(
                worksheet.cell(
                    row=1,
                    column=column,
                ).value
                or ""
            ).strip(): column
            for column in range(
                1,
                worksheet.max_column + 1,
            )
        }

        missing = (
            set(self.HEADERS)
            - set(headers)
        )

        if missing:
            raise ValueError(
                "Faltan columnas en revision: "
                f"{sorted(missing)}"
            )

        resolved = []
        pending = []
        invalid = []
        seen_ids = set()
        rows = []

        for row_number in range(
            2,
            worksheet.max_row + 1,
        ):
            item_id = self._normalize_item_id(
                worksheet.cell(
                    row=row_number,
                    column=headers[
                        "ITEM ID"
                    ],
                ).value
            )

            if not item_id:
                continue

            email = str(
                worksheet.cell(
                    row=row_number,
                    column=headers[
                        "CORREO"
                    ],
                ).value
                or ""
            ).strip().lower()

            nit = str(
                worksheet.cell(
                    row=row_number,
                    column=headers[
                        "NIT"
                    ],
                ).value
                or ""
            ).strip()

            received = str(
                worksheet.cell(
                    row=row_number,
                    column=headers[
                        "CLIENTE RECIBIDO"
                    ],
                ).value
                or ""
            ).strip()

            assigned = str(
                worksheet.cell(
                    row=row_number,
                    column=headers[
                        "CLIENTE ASIGNADO"
                    ],
                ).value
                or ""
            ).strip()

            base = {
                "_item_id": item_id,
                "email": email,
                "nit": nit,
                "cliente_recibido":
                    received,
                "cliente_asignado":
                    assigned,
                "row_number":
                    row_number,
            }

            if item_id in seen_ids:
                invalid.append({
                    **base,
                    "error":
                        "ITEM ID duplicado "
                        "en revision",
                })
                continue

            seen_ids.add(
                item_id
            )

            if not assigned:
                pending.append(
                    base
                )
                rows.append({
                    **base,
                    "estado":
                        "PENDIENTE_CLIENTE",
                })
                continue

            canonical = (
                self.config.resolve_alias(
                    country_code,
                    assigned,
                )
            )

            if not canonical:
                invalid.append({
                    **base,
                    "error":
                        "Cliente asignado no "
                        "pertenece al catalogo",
                })
                rows.append({
                    **base,
                    "estado":
                        "CLIENTE_INVALIDO",
                })
                continue

            resolved_row = {
                **base,
                "cliente_canonico":
                    canonical,
            }

            resolved.append(
                resolved_row
            )

            rows.append({
                **resolved_row,
                "estado":
                    "MANUAL_VALIDADO",
            })

        return {
            "resolved": resolved,
            "pending": pending,
            "invalid": invalid,
            "rows": rows,
        }

    def _read_existing_assignments(
        self,
        output_path: Path,
    ) -> dict:
        if not output_path.exists():
            return {}

        try:
            workbook = load_workbook(
                output_path,
                data_only=False,
            )

            if (
                "REVISION_CLIENTES"
                not in workbook.sheetnames
            ):
                return {}

            worksheet = workbook[
                "REVISION_CLIENTES"
            ]

            header_map = {
                str(
                    worksheet.cell(
                        row=1,
                        column=column,
                    ).value
                    or ""
                ).strip(): column
                for column in range(
                    1,
                    worksheet.max_column + 1,
                )
            }

            item_col = header_map.get(
                "ITEM ID"
            )

            assigned_col = (
                header_map.get(
                    "CLIENTE ASIGNADO"
                )
            )

            if (
                not item_col
                or not assigned_col
            ):
                return {}

            assignments = {}

            for row in range(
                2,
                worksheet.max_row + 1,
            ):
                item_id = (
                    self._normalize_item_id(
                        worksheet.cell(
                            row=row,
                            column=item_col,
                        ).value
                    )
                )

                assigned = str(
                    worksheet.cell(
                        row=row,
                        column=assigned_col,
                    ).value
                    or ""
                ).strip()

                if item_id and assigned:
                    assignments[
                        item_id
                    ] = assigned

            return assignments

        except Exception:
            return {}

    @staticmethod
    def _normalize_item_id(
        value,
    ) -> str:
        text = str(
            value or ""
        ).strip()

        if text.endswith(".0"):
            prefix = text[:-2]

            if prefix.isdigit():
                return prefix

        return text
