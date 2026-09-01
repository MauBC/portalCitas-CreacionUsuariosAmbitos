from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.dataframe import dataframe_to_rows


class ControlSheetService:

    SHEET_NAME = "CONTROL"

    @staticmethod
    def _as_bool(value) -> bool:

        if value is True:
            return True

        if value is False or value is None:
            return False

        try:
            if pd.isna(value):
                return False
        except (TypeError, ValueError):
            pass

        return str(value).strip().lower() in {
            "1",
            "true",
            "si",
            "sí",
            "yes",
        }

    @staticmethod
    def _text(value) -> str:

        if value is None:
            return ""

        try:
            if pd.isna(value):
                return ""
        except (TypeError, ValueError):
            pass

        return str(value).strip()

    def build_control(
        self,
        valid_path: str | Path,
        error_path: str | Path,
    ) -> pd.DataFrame:

        valid = pd.read_excel(
            valid_path,
            sheet_name="DATOS_TECNICOS",
            dtype=str,
        )

        errors = pd.read_excel(
            error_path,
            sheet_name="DATOS_TECNICOS",
            dtype=str,
        )

        rows = []

        # ==========================================
        # VALIDOS
        # ==========================================

        for _, row in valid.iterrows():

            cliente_original = self._text(
                row.get("nombre_cliente")
            )

            cliente_normalizado = self._text(
                row.get("cliente_normalizado")
            )

            if not cliente_normalizado:
                cliente_normalizado = cliente_original

            rows.append({
                "ITEM_ID":
                    self._text(row.get("_item_id")),

                "EMAIL":
                    self._text(row.get("email")),

                "PROVEEDOR":
                    self._text(
                        row.get("nombre_proveedor")
                    ),

                "NIT_PROVEEDOR":
                    self._text(
                        row.get("nit_proveedor")
                    ),

                "CLIENTE_ORIGINAL":
                    cliente_original,

                "CLIENTE_NORMALIZADO":
                    cliente_normalizado,

                "ESTADO":
                    "LISTO_PARA_CARGA",

                "OBSERVACION":
                    "",
            })

        # ==========================================
        # ERRORES
        # ==========================================

        for _, row in errors.iterrows():

            revision_cliente = self._as_bool(
                row.get(
                    "requiere_revision_cliente"
                )
            )

            estado = (
                "REVISION_CLIENTE"
                if revision_cliente
                else "ERROR_DATOS"
            )

            cliente_normalizado = self._text(
                row.get("cliente_normalizado")
            )

            # Si requiere revision no mostramos
            # un candidato como si ya fuese correcto.
            if revision_cliente:
                cliente_normalizado = ""

            rows.append({
                "ITEM_ID":
                    self._text(row.get("_item_id")),

                "EMAIL":
                    self._text(row.get("email")),

                "PROVEEDOR":
                    self._text(
                        row.get("nombre_proveedor")
                    ),

                "NIT_PROVEEDOR":
                    self._text(
                        row.get("nit_proveedor")
                    ),

                "CLIENTE_ORIGINAL":
                    self._text(
                        row.get("nombre_cliente")
                    ),

                "CLIENTE_NORMALIZADO":
                    cliente_normalizado,

                "ESTADO":
                    estado,

                "OBSERVACION":
                    self._text(
                        row.get("observaciones")
                    ),
            })

        return pd.DataFrame(rows)

    def add_to_template(
        self,
        template_path: str | Path,
        valid_path: str | Path,
        error_path: str | Path,
    ) -> pd.DataFrame:

        template_path = Path(template_path)

        control = self.build_control(
            valid_path=valid_path,
            error_path=error_path,
        )

        wb = load_workbook(
            template_path
        )

        if self.SHEET_NAME in wb.sheetnames:
            del wb[self.SHEET_NAME]

        ws = wb.create_sheet(
            self.SHEET_NAME
        )

        for row in dataframe_to_rows(
            control,
            index=False,
            header=True,
        ):
            ws.append(row)

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        widths = {
            "A": 12,
            "B": 38,
            "C": 45,
            "D": 20,
            "E": 45,
            "F": 25,
            "G": 24,
            "H": 80,
        }

        for column, width in widths.items():
            ws.column_dimensions[column].width = width

        wb.save(
            template_path
        )

        return control
