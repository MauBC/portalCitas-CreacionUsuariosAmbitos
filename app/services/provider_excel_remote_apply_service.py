from __future__ import annotations

from datetime import datetime
from hashlib import sha256
from io import BytesIO
from pathlib import Path, PurePosixPath
import re
from zipfile import ZipFile
from xml.etree import ElementTree as ET

import pandas as pd
import requests
from openpyxl import load_workbook

from app.auth.msal_auth import MsalAuthService
from app.services.provider_excel_remote_status_service import (
    ProviderExcelRemoteStatusService,
)


class ProviderExcelRemoteApplyService:
    WORKBOOK_NS = (
        "http://schemas.openxmlformats.org/"
        "spreadsheetml/2006/main"
    )
    OFFICE_REL_NS = (
        "http://schemas.openxmlformats.org/"
        "officeDocument/2006/relationships"
    )
    PACKAGE_REL_NS = (
        "http://schemas.openxmlformats.org/"
        "package/2006/relationships"
    )

    def __init__(
        self,
        status_service: ProviderExcelRemoteStatusService | None = None,
    ):
        self.status_service = (
            status_service
            or ProviderExcelRemoteStatusService()
        )

    def resolve_drive_item(
        self,
        shared_url: str,
    ) -> dict:
        url = str(
            shared_url or ""
        ).strip()

        if not url:
            raise ValueError(
                "Falta URL compartida de SharePoint/OneDrive."
            )

        encoded = (
            self.status_service
            ._encode_shared_url(
                url
            )
        )

        token = (
            MsalAuthService()
            .get_access_token()
        )

        endpoint = (
            "https://graph.microsoft.com/v1.0/"
            f"shares/{encoded}/driveItem"
        )

        with requests.get(
            endpoint,
            headers={
                "Authorization":
                    f"Bearer {token}",
            },
            timeout=120,
        ) as response:
            if not response.ok:
                detail = (
                    response.text[:1000]
                    if response.text
                    else ""
                )

                raise RuntimeError(
                    "No se pudo resolver el archivo remoto. "
                    f"HTTP {response.status_code}. "
                    f"{detail}"
                )

            data = response.json()

        item_id = str(
            data.get("id")
            or ""
        ).strip()

        parent = (
            data.get("parentReference")
            or {}
        )

        drive_id = str(
            parent.get("driveId")
            or ""
        ).strip()

        if not item_id or not drive_id:
            raise RuntimeError(
                "Graph no devolvio driveId/itemId "
                "del archivo compartido."
            )

        return {
            "item_id": item_id,
            "drive_id": drive_id,
            "name": str(
                data.get("name")
                or ""
            ).strip(),
            "etag": str(
                data.get("eTag")
                or ""
            ).strip(),
            "size": data.get("size"),
        }

    def upload_content(
        self,
        metadata: dict,
        content: bytes,
    ) -> dict:
        if not content:
            raise ValueError(
                "El contenido a subir esta vacio."
            )

        drive_id = str(
            metadata.get("drive_id")
            or ""
        ).strip()

        item_id = str(
            metadata.get("item_id")
            or ""
        ).strip()

        if not drive_id or not item_id:
            raise ValueError(
                "Faltan drive_id/item_id para subir el archivo."
            )

        token = (
            MsalAuthService()
            .get_access_token()
        )

        endpoint = (
            "https://graph.microsoft.com/v1.0/"
            f"drives/{drive_id}/items/{item_id}/content"
        )

        with requests.put(
            endpoint,
            headers={
                "Authorization":
                    f"Bearer {token}",
                "Content-Type":
                    "application/vnd.openxmlformats-"
                    "officedocument.spreadsheetml.sheet",
            },
            data=content,
            timeout=120,
        ) as response:
            if not response.ok:
                detail = (
                    response.text[:1000]
                    if response.text
                    else ""
                )

                raise RuntimeError(
                    "No se pudo reemplazar el Excel remoto. "
                    f"HTTP {response.status_code}. "
                    f"{detail}"
                )

            return response.json()

    def build_updated_content(
        self,
        content: bytes,
        plan: pd.DataFrame,
    ) -> bytes:
        if not content:
            raise ValueError(
                "El Excel remoto esta vacio."
            )

        required = {
            "FILA EXCEL",
            "CREADO NUEVO",
            "ACCION",
        }

        missing = required - set(
            plan.columns
        )

        if missing:
            raise ValueError(
                "Faltan columnas en plan Fase 2: "
                f"{sorted(missing)}"
            )

        update_rows = plan[
            plan["ACCION"]
            .fillna("")
            .astype(str)
            .eq("ACTUALIZAR")
        ].copy()

        if update_rows.empty:
            return content

        workbook = load_workbook(
            BytesIO(content),
            read_only=True,
            data_only=False,
        )

        try:
            if self.status_service.SHEET_NAME not in (
                workbook.sheetnames
            ):
                raise ValueError(
                    "El Excel remoto no contiene la hoja "
                    f"'{self.status_service.SHEET_NAME}'."
                )

            worksheet = workbook[
                self.status_service.SHEET_NAME
            ]

            headers = (
                self.status_service
                ._build_header_map(
                    worksheet
                )
            )

            creado_key = (
                self.status_service
                ._normalize_header(
                    "creado"
                )
            )

            if creado_key not in headers:
                raise ValueError(
                    "El Excel remoto no contiene "
                    "la columna creado."
                )

            creado_column = headers[
                creado_key
            ]

        finally:
            workbook.close()

        column_letter = (
            self._column_letter(
                creado_column
            )
        )

        sheet_path = (
            self._find_sheet_xml_path(
                content,
                self.status_service.SHEET_NAME,
            )
        )

        with ZipFile(
            BytesIO(content),
            "r",
        ) as source_zip:
            sheet_xml = source_zip.read(
                sheet_path
            )

            patched_xml = sheet_xml

            for _, row in (
                update_rows.iterrows()
            ):
                excel_row = self._to_int(
                    row.get("FILA EXCEL")
                )

                desired = self._to_int(
                    row.get("CREADO NUEVO")
                )

                if (
                    excel_row is None
                    or excel_row < 2
                ):
                    raise ValueError(
                        "Fila Excel invalida en plan Fase 2."
                    )

                if desired not in {
                    1,
                    2,
                }:
                    raise ValueError(
                        "CREADO NUEVO debe ser 1 o 2."
                    )

                cell_ref = (
                    f"{column_letter}{excel_row}"
                )

                patched_xml = (
                    self._patch_numeric_cell(
                        xml=patched_xml,
                        row_number=excel_row,
                        cell_ref=cell_ref,
                        value=desired,
                    )
                )

            output = BytesIO()

            from zipfile import ZipFile as OutputZip

            with OutputZip(
                output,
                "w",
            ) as target_zip:
                for info in (
                    source_zip.infolist()
                ):
                    data = (
                        patched_xml
                        if info.filename
                        == sheet_path
                        else source_zip.read(
                            info.filename
                        )
                    )

                    target_zip.writestr(
                        info,
                        data,
                    )

        updated = output.getvalue()

        self._assert_other_zip_members_unchanged(
            original=content,
            updated=updated,
            changed_member=sheet_path,
        )

        return updated

    def apply(
        self,
        shared_url: str,
        relation_results: pd.DataFrame,
        output_dir: str,
    ) -> dict:
        folder = Path(
            output_dir
        )

        folder.mkdir(
            parents=True,
            exist_ok=True,
        )

        timestamp = datetime.now().strftime(
            "%Y%m%d_%H%M%S"
        )

        metadata_before = (
            self.resolve_drive_item(
                shared_url
            )
        )

        original = (
            self.status_service
            .download(
                shared_url
            )
        )

        preview = (
            self.status_service
            .build_preview(
                content=original,
                relation_results=
                    relation_results,
            )
        )

        if not preview["ready"]:
            raise RuntimeError(
                "Fase 2 no esta lista para escritura. "
                "Ejecuta PREVIEW y corrige los bloqueos."
            )

        updated = self.build_updated_content(
            content=original,
            plan=preview["plan"],
        )

        local_verify = (
            self.status_service
            .build_preview(
                content=updated,
                relation_results=
                    relation_results,
            )
        )

        self._assert_fully_applied(
            verification=local_verify,
            expected_rows=len(
                relation_results
            ),
        )

        backup_path = (
            folder
            / (
                "excel_sharepoint_BACKUP_ANTES_FASE2_"
                f"{timestamp}.xlsx"
            )
        )

        candidate_path = (
            folder
            / (
                "excel_sharepoint_OBJETIVO_FASE2_"
                f"{timestamp}.xlsx"
            )
        )

        backup_path.write_bytes(
            original
        )

        candidate_path.write_bytes(
            updated
        )

        metadata_now = (
            self.resolve_drive_item(
                shared_url
            )
        )

        if (
            metadata_before["item_id"]
            != metadata_now["item_id"]
            or metadata_before["drive_id"]
            != metadata_now["drive_id"]
        ):
            raise RuntimeError(
                "El archivo compartido cambio de identidad "
                "antes de la escritura."
            )

        before_etag = str(
            metadata_before.get("etag")
            or ""
        )

        now_etag = str(
            metadata_now.get("etag")
            or ""
        )

        if (
            before_etag
            and now_etag
            and before_etag != now_etag
        ):
            raise RuntimeError(
                "El Excel remoto fue modificado entre la "
                "validacion y la escritura. No se subio nada."
            )

        upload_result = self.upload_content(
            metadata=metadata_now,
            content=updated,
        )

        uploaded_etag = str(
            upload_result.get("eTag")
            or ""
        )

        remote_after = (
            self.status_service
            .download(
                shared_url
            )
        )

        verification = (
            self.status_service
            .build_preview(
                content=remote_after,
                relation_results=
                    relation_results,
            )
        )

        try:
            self._assert_fully_applied(
                verification=verification,
                expected_rows=len(
                    relation_results
                ),
            )
        except Exception:
            rollback = (
                self._try_rollback(
                    shared_url=shared_url,
                    original=original,
                    metadata=metadata_now,
                    uploaded_etag=uploaded_etag,
                )
            )

            if rollback:
                raise RuntimeError(
                    "La verificacion posterior fallo y se "
                    "restauro automaticamente el backup."
                )

            raise RuntimeError(
                "La verificacion posterior fallo. "
                f"Backup local: {backup_path}"
            )

        verified_path = (
            folder
            / (
                "excel_sharepoint_VERIFICADO_DESPUES_FASE2_"
                f"{timestamp}.xlsx"
            )
        )

        verified_path.write_bytes(
            remote_after
        )

        plan = verification[
            "plan"
        ].copy()

        status_numeric = pd.to_numeric(
            plan["CREADO ACTUAL"],
            errors="coerce",
        )

        created_1 = int(
            status_numeric.eq(1).sum()
        )

        created_2 = int(
            status_numeric.eq(2).sum()
        )

        report_path = (
            folder
            / (
                "fase2_PER_APLICACION_SHAREPOINT_"
                f"{timestamp}.xlsx"
            )
        )

        summary = {
            "filas_verificadas": len(plan),
            "creado_1": created_1,
            "creado_2": created_2,
            "sin_cambio_post": int(
                verification["summary"][
                    "sin_cambio"
                ]
            ),
            "errores_verificacion_post": int(
                verification["summary"][
                    "errores_verificacion"
                ]
            ),
            "conflictos_post": int(
                verification["summary"][
                    "conflictos"
                ]
            ),
            "estados_invalidos_post": int(
                verification["summary"][
                    "estados_invalidos"
                ]
            ),
            "sha256_original": sha256(
                original
            ).hexdigest(),
            "sha256_subido": sha256(
                updated
            ).hexdigest(),
            "sha256_releido": sha256(
                remote_after
            ).hexdigest(),
            "backup_path": str(
                backup_path.resolve()
            ),
            "candidate_path": str(
                candidate_path.resolve()
            ),
            "verified_path": str(
                verified_path.resolve()
            ),
        }

        with pd.ExcelWriter(
            report_path,
            engine="openpyxl",
        ) as writer:
            pd.DataFrame([
                {
                    "METRICA": key,
                    "VALOR": value,
                }
                for key, value
                in summary.items()
            ]).to_excel(
                writer,
                sheet_name="RESUMEN",
                index=False,
            )

            plan.to_excel(
                writer,
                sheet_name="VERIFICACION",
                index=False,
            )

        return {
            "ok": True,
            "created_1": created_1,
            "created_2": created_2,
            "verified_rows": len(plan),
            "backup_path": str(
                backup_path.resolve()
            ),
            "candidate_path": str(
                candidate_path.resolve()
            ),
            "verified_path": str(
                verified_path.resolve()
            ),
            "report_path": str(
                report_path.resolve()
            ),
            "sha256_uploaded": sha256(
                updated
            ).hexdigest(),
            "sha256_downloaded": sha256(
                remote_after
            ).hexdigest(),
        }

    def _try_rollback(
        self,
        shared_url: str,
        original: bytes,
        metadata: dict,
        uploaded_etag: str,
    ) -> bool:
        try:
            current = self.resolve_drive_item(
                shared_url
            )

            current_etag = str(
                current.get("etag")
                or ""
            )

            if (
                uploaded_etag
                and current_etag
                and uploaded_etag
                != current_etag
            ):
                return False

            self.upload_content(
                metadata=current,
                content=original,
            )

            restored = (
                self.status_service
                .download(
                    shared_url
                )
            )

            return (
                sha256(restored).hexdigest()
                == sha256(original).hexdigest()
            )

        except Exception:
            return False

    @staticmethod
    def _assert_fully_applied(
        verification: dict,
        expected_rows: int,
    ) -> None:
        summary = verification[
            "summary"
        ]

        if not verification["ready"]:
            raise RuntimeError(
                "El Excel no paso la verificacion Fase 2."
            )

        if int(
            summary["actualizaciones"]
        ) != 0:
            raise RuntimeError(
                "Todavia quedan cambios Fase 2 pendientes."
            )

        if int(
            summary["sin_cambio"]
        ) != int(
            expected_rows
        ):
            raise RuntimeError(
                "No todas las filas quedaron con el estado esperado."
            )

    @staticmethod
    def _column_letter(
        column_number: int,
    ) -> str:
        if column_number <= 0:
            raise ValueError(
                "Numero de columna invalido."
            )

        result = ""
        value = column_number

        while value:
            value, remainder = divmod(
                value - 1,
                26,
            )
            result = (
                chr(65 + remainder)
                + result
            )

        return result

    @classmethod
    def _find_sheet_xml_path(
        cls,
        content: bytes,
        sheet_name: str,
    ) -> str:
        with ZipFile(
            BytesIO(content),
            "r",
        ) as archive:
            workbook_xml = ET.fromstring(
                archive.read(
                    "xl/workbook.xml"
                )
            )

            relationship_id = None

            for sheet in workbook_xml.findall(
                f".//{{{cls.WORKBOOK_NS}}}sheet"
            ):
                if sheet.attrib.get(
                    "name"
                ) == sheet_name:
                    relationship_id = sheet.attrib.get(
                        f"{{{cls.OFFICE_REL_NS}}}id"
                    )
                    break

            if not relationship_id:
                raise ValueError(
                    f"No se encontro hoja {sheet_name}."
                )

            rels_xml = ET.fromstring(
                archive.read(
                    "xl/_rels/workbook.xml.rels"
                )
            )

            target = None

            for relation in rels_xml.findall(
                f"{{{cls.PACKAGE_REL_NS}}}Relationship"
            ):
                if relation.attrib.get(
                    "Id"
                ) == relationship_id:
                    target = relation.attrib.get(
                        "Target"
                    )
                    break

            if not target:
                raise RuntimeError(
                    "No se pudo resolver XML de la hoja."
                )

        if target.startswith("/"):
            return target.lstrip("/")

        return str(
            PurePosixPath("xl")
            / PurePosixPath(target)
        )

    @staticmethod
    def _patch_numeric_cell(
        xml: bytes,
        row_number: int,
        cell_ref: str,
        value: int,
    ) -> bytes:
        text = xml.decode("utf-8")

        escaped_ref = re.escape(
            cell_ref
        )

        full_pattern = re.compile(
            rf'<c\b(?=[^>]*\br="{escaped_ref}")'
            rf'(?P<attrs>[^>]*)>(?P<body>.*?)</c>',
            re.DOTALL,
        )

        match = full_pattern.search(
            text
        )

        if match:
            attrs = re.sub(
                r'\s+t="[^"]*"',
                "",
                match.group("attrs"),
            )

            replacement = (
                f'<c{attrs}><v>{value}</v></c>'
            )

            return (
                text[:match.start()]
                + replacement
                + text[match.end():]
            ).encode("utf-8")

        self_closing = re.compile(
            rf'<c\b(?=[^>]*\br="{escaped_ref}")'
            rf'(?P<attrs>[^>]*)/>',
            re.DOTALL,
        )

        match = self_closing.search(
            text
        )

        if match:
            attrs = re.sub(
                r'\s+t="[^"]*"',
                "",
                match.group("attrs"),
            )

            replacement = (
                f'<c{attrs}><v>{value}</v></c>'
            )

            return (
                text[:match.start()]
                + replacement
                + text[match.end():]
            ).encode("utf-8")

        row_pattern = re.compile(
            rf'(<row\b(?=[^>]*\br="{row_number}")'
            rf'[^>]*>)(?P<body>.*?)(</row>)',
            re.DOTALL,
        )

        row_match = row_pattern.search(
            text
        )

        if not row_match:
            raise RuntimeError(
                f"No se encontro fila {row_number} "
                "en XML de Proveedores."
            )

        new_cell = (
            f'<c r="{cell_ref}"><v>{value}</v></c>'
        )

        replacement = (
            row_match.group(1)
            + row_match.group("body")
            + new_cell
            + row_match.group(3)
        )

        return (
            text[:row_match.start()]
            + replacement
            + text[row_match.end():]
        ).encode("utf-8")

    @staticmethod
    def _assert_other_zip_members_unchanged(
        original: bytes,
        updated: bytes,
        changed_member: str,
    ) -> None:
        with ZipFile(
            BytesIO(original),
            "r",
        ) as original_zip:
            with ZipFile(
                BytesIO(updated),
                "r",
            ) as updated_zip:
                original_names = set(
                    original_zip.namelist()
                )

                updated_names = set(
                    updated_zip.namelist()
                )

                if original_names != updated_names:
                    raise RuntimeError(
                        "La estructura interna del XLSX cambio."
                    )

                for name in original_names:
                    if name == changed_member:
                        continue

                    if (
                        original_zip.read(name)
                        != updated_zip.read(name)
                    ):
                        raise RuntimeError(
                            "Se modifico un componente XLSX "
                            f"no permitido: {name}"
                        )

    @staticmethod
    def _to_int(
        value,
    ):
        try:
            numeric = float(
                str(value).strip()
            )
        except (
            TypeError,
            ValueError,
        ):
            return None

        if not numeric.is_integer():
            return None

        return int(numeric)
