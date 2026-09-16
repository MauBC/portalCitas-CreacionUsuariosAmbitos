import os
import shutil
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from app.config.paths import PROJECT_ROOT, TEMPLATES_ROOT


class AmbitosExcelService:

    SHEET_CONFIG = {
        "Entidades": [
            "Razon Social",
            "Apellidos y nombres",
            "Tipo",
            "RUC",
            "Email",
            "Telefono",
            "Direccion",
            "Tipo Usuario",
            "Tipo de negocio Nuevo",
            "Tipo de negocio Eliminar",
            "Relación Nueva",
            "Relacion Eliminada",
            "Estado",
            "Comentario",
        ],
        "Ambitos": [
            "Email",
            "Empresa",
            "Todos los negocios",
            "Todas las sedes",
            "Sedes Nuevas",
            "Sedes Eliminadas",
            "Estado",
            "Comentario",
        ],
        "RelacionNueva": [
            "Grupo",
            "Razon Social",
            "RUC",
        ],
        "TipoNegocioNuevo": [
            "Grupo",
            "Tipo de negocio",
            "Rubro",
        ],
        "SedeNueva": [
            "Grupo",
            "Sede",
            "Tipo de negocio",
        ],
    }

    DATA_KEYS = {
        "Entidades": "entidades",
        "Ambitos": "ambitos",
        "RelacionNueva": "relacion_nueva",
        "TipoNegocioNuevo":
            "tipo_negocio_nuevo",
        "SedeNueva": "sede_nueva",
    }

    def __init__(
        self,
        output_dir: str = "salidas",
        template_path: str | None = None,
    ):
        self.output_dir = output_dir

        os.makedirs(
            self.output_dir,
            exist_ok=True,
        )

        configured = (
            Path(template_path)
            if template_path
            else (
                TEMPLATES_ROOT
                / "subida_ambitos.xlsx"
            )
        )

        if not configured.exists():
            root_candidate = (
                PROJECT_ROOT
                / "subida_ambitos.xlsx"
            )

            if root_candidate.exists():
                configured = root_candidate

        self.template_path = configured

    def export_template_ambitos(
        self,
        data: dict | None = None,
        output_path: str | None = None,
    ) -> str:
        data = data or {}

        if not self.template_path.exists():
            raise FileNotFoundError(
                "No existe la plantilla oficial "
                "subida_ambitos.xlsx. "
                f"Ruta esperada: "
                f"{self.template_path}"
            )

        if not output_path:
            timestamp = (
                datetime.now()
                .strftime("%Y%m%d_%H%M%S")
            )

            output_path = os.path.join(
                self.output_dir,
                (
                    "plantilla_ambitos_"
                    f"{timestamp}.xlsx"
                ),
            )

        output_path = Path(
            output_path
        )

        output_path.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        shutil.copy2(
            self.template_path,
            output_path,
        )

        workbook = load_workbook(
            output_path
        )

        self._validate_template(
            workbook
        )

        for sheet_name, headers in (
            self.SHEET_CONFIG.items()
        ):
            data_key = self.DATA_KEYS[
                sheet_name
            ]

            rows = data.get(
                data_key,
                [],
            )

            self._replace_rows(
                worksheet=workbook[
                    sheet_name
                ],
                headers=headers,
                rows=rows,
            )

        workbook.save(
            output_path
        )

        return str(
            output_path.resolve()
        )

    def _validate_template(
        self,
        workbook,
    ):
        expected_sheets = {
            "Entidades",
            "Ambitos",
            "RelacionNueva",
            "RelacionEliminar",
            "TipoNegocioNuevo",
            "TipoNegocioEliminar",
            "SedeNueva",
            "SedeEliminar",
        }

        missing = (
            expected_sheets
            - set(workbook.sheetnames)
        )

        if missing:
            raise ValueError(
                "La plantilla de ambitos "
                "no tiene todas las hojas: "
                f"{sorted(missing)}"
            )

        for sheet_name, headers in (
            self.SHEET_CONFIG.items()
        ):
            worksheet = workbook[
                sheet_name
            ]

            current_headers = [
                worksheet.cell(
                    row=1,
                    column=index,
                ).value
                for index in range(
                    1,
                    len(headers) + 1,
                )
            ]

            if current_headers != headers:
                raise ValueError(
                    "Encabezados inesperados "
                    f"en {sheet_name}. "
                    f"Esperado: {headers}. "
                    f"Actual: {current_headers}"
                )

    def _replace_rows(
        self,
        worksheet,
        headers: list[str],
        rows: list[dict],
    ):
        if worksheet.max_row > 1:
            worksheet.delete_rows(
                2,
                worksheet.max_row - 1,
            )

        for row_number, row in enumerate(
            rows,
            start=2,
        ):
            for column, header in enumerate(
                headers,
                start=1,
            ):
                value = self._get_value(
                    row,
                    header,
                )

                worksheet.cell(
                    row=row_number,
                    column=column,
                    value=value,
                )

    @staticmethod
    def _get_value(
        row: dict,
        header: str,
    ):
        if header == "Relación Nueva":
            value = (
                row.get(
                    "Relación Nueva"
                )
                if (
                    "Relación Nueva"
                    in row
                )
                else row.get(
                    "Relacion Nueva",
                    "",
                )
            )
        else:
            value = row.get(
                header,
                "",
            )

        if header in {
            "Todos los negocios",
            "Todas las sedes",
        }:
            if isinstance(
                value,
                bool,
            ):
                return str(
                    value
                ).lower()

        if header in {
            "RUC",
            "Empresa",
        }:
            return str(
                value or ""
            )

        return (
            ""
            if value is None
            else value
        )
